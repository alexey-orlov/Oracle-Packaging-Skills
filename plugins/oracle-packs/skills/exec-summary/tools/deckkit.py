"""SoftServe deck primitives shared by the deck and exec-summary builders.

Brand tokens, shape/text helpers, the pack-spec accessor and the headless
text-fit estimator. No machine-specific paths: fonts are probed at run time
from a candidate list — the brand face the plugin ships in its fonts/ folder
first, then Helvetica-metric stand-ins — and the deck base ships next to the
skill.

This file is duplicated verbatim in:
    skills/deck/tools/deckkit.py
    skills/exec-summary/tools/deckkit.py
so each skill stays copy-standalone. Keep the two in sync — the builders
prefer the copy sitting next to them and fall back to the sibling skill's.
"""

from __future__ import annotations

import os
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

EMU_IN = 914400


def inch(v: float) -> int:
    return int(round(v * EMU_IN))


# --------------------------------------------------------------------------
# Brand tokens (measured from the SoftServe EMEA master and the WfO artifacts)
# --------------------------------------------------------------------------

C = {
    # ink / neutrals
    "ink": "26282B",          # body black, dark bands
    "ink_soft": "44515A",     # secondary ink, neutral accent bars
    "muted": "6B7076",        # secondary body text
    "muted_light": "8A9095",  # footnote / source text
    "hairline": "DCE1E5",     # panel outlines
    "hairline_alt": "D2D8DD",
    "panel_grey": "F4F6F7",   # neutral panel fill
    "white": "FFFFFF",
    # blue (Oracle / solution side)
    "blue": "1485C3",
    "blue_mid": "459FDD",
    "blue_light": "6DB2E2",
    "blue_dark": "0E5E8A",
    "blue_tint": "EAF3FB",    # panel tint
    "blue_tint_2": "ECF5FB",
    # orange (SoftServe accent / proof / attention)
    "orange": "F36949",
    "orange_light": "FE8D6B",  # tier S header
    "orange_dark": "D84F2E",   # tier L header
    "orange_tint": "FDEDE8",
}

FONT_BODY = "Replica LL TT"     # the master's body face
FONT_MONO = "Roboto Mono"       # small keys / numerals
FONT_TITLE = "+mj-lt"           # theme major latin (Azurio) — never hardcode

# Canvas + master furniture (inches), 13.33 x 7.5 in
CANVAS_W, CANVAS_H = 13.333, 7.5
MARGIN_L = 0.42
CONTENT_W = 12.49
HEADER_BOX = (6.79, 0.31, 5.48, 0.23)      # running header placeholder
TITLE_BOX = (0.39, 1.18, 12.05, 0.43)      # Title-1Column title placeholder
DECK_TITLE_BOX = (0.39, 1.40, 11.80, 0.95)  # sales-deck content-slide title
FOOTNOTE_Y = 6.80


# --------------------------------------------------------------------------
# Fonts for the headless fit estimate: the shipped brand face, else stand-ins
# --------------------------------------------------------------------------

# The brand face ships privately in the plugin's own fonts/ folder (<plugin>/fonts, three levels
# above this tools/ folder), for practice members; measured with it, the estimate is the real one.
# A folder with no files in it is passed over quietly and the stand-ins below take over.
PLUGIN_FONTS = Path(__file__).resolve().parents[3] / "fonts"
_SHIPPED_FACES = {
    False: str(PLUGIN_FONTS / "ReplicaLLTT-Regular.ttf"),
    True: str(PLUGIN_FONTS / "ReplicaLLTT-Bold.ttf"),
}

_FONT_CANDIDATES = {
    False: [  # regular
        _SHIPPED_FACES[False],
        "/Applications/LibreOffice.app/Contents/Resources/fonts/truetype/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ],
    True: [  # bold
        _SHIPPED_FACES[True],
        "/Applications/LibreOffice.app/Contents/Resources/fonts/truetype/LiberationSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
        "/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ],
}

# The stand-ins are Helvetica-metric, which tracks Replica closely; SAFETY covers the residual
# difference. Measured on the shipped Replica itself there is no metric difference left to cover,
# only rendering slack, so the margin drops to SAFETY_SHIPPED. Rule: trust geometry, not glyphs.
SAFETY = 1.06          # +6% width headroom when a stand-in is measured
SAFETY_SHIPPED = 1.02  # +2% when the shipped brand face is measured
LINE_FACTOR = 1.22     # single-spaced line box as a multiple of the point size

_font_cache: dict[tuple[float, bool], Any] = {}
_font_source: dict[bool, str] = {}   # bold -> the file measured, or a note when there is none


def _load_metric_font(pt: float, bold: bool):
    key = (round(pt, 2), bold)
    if key in _font_cache:
        return _font_cache[key]
    try:
        from PIL import ImageFont
    except ImportError:  # pragma: no cover - Pillow is a declared dependency
        _font_cache[key] = None
        return None
    size = max(1, int(round(pt * 20)))  # measure at 20x for sub-point accuracy
    for path in _FONT_CANDIDATES[bold]:
        if not os.path.exists(path):
            continue
        try:
            f = ImageFont.truetype(path, size)
        except Exception:
            continue
        _font_source.setdefault(bold, path)
        _font_cache[key] = f
        return f
    try:
        from PIL import ImageFont
        f = ImageFont.load_default()
    except Exception:
        f = None
    _font_source.setdefault(bold, "PIL default bitmap font (crude estimate)")
    _font_cache[key] = f
    return f


def fit_safety(bold: bool = False) -> float:
    """The width margin on the face measured for this weight: SAFETY_SHIPPED on the shipped brand
    face, SAFETY on a stand-in."""
    if bold not in _font_source:
        _load_metric_font(10, bold)
    return SAFETY_SHIPPED if _font_source.get(bold) == _SHIPPED_FACES[bold] else SAFETY


def metric_font_source() -> str:
    """The file the estimate measures regular text with."""
    if False not in _font_source:
        _load_metric_font(10, False)
    return _font_source.get(False, "none")


def metric_font_line() -> str:
    """The fit report's first line: which face was measured, where from, and the margin on it."""
    source = metric_font_source()
    margin = int(round((fit_safety(False) - 1) * 100))
    if source == _SHIPPED_FACES[False]:
        return f"metric: {source} (the shipped brand face)  (+{margin}% safety)"
    return f"metric stand-in: {source}  (+{margin}% safety)"


def text_width_in(s: str, pt: float, bold: bool = False) -> float:
    """Width of `s` in inches at `pt`, measured on the shipped face or a stand-in, +safety."""
    if not s:
        return 0.0
    f = _load_metric_font(pt, bold)
    margin = fit_safety(bold)
    if f is None:
        return len(s) * pt * 0.5 / 72.0 * margin
    try:
        w = f.getlength(s) / 20.0
    except AttributeError:  # very old Pillow / default bitmap font
        w = f.getsize(s)[0] / 20.0
    return w / 72.0 * margin


def wrap_count(s: str, pt: float, box_w_in: float, bold: bool = False) -> int:
    """Greedy word-wrap line count for `s` inside `box_w_in`."""
    if not s.strip():
        return 1
    if box_w_in <= 0:
        return 999
    lines, cur = 1, ""
    for word in s.split():
        trial = word if not cur else cur + " " + word
        if text_width_in(trial, pt, bold) <= box_w_in:
            cur = trial
        else:
            if cur:
                lines += 1
            cur = word
            while text_width_in(cur, pt, bold) > box_w_in and len(cur) > 1:
                # a single unbreakable token wider than the box
                cut = max(1, int(len(cur) * box_w_in / max(text_width_in(cur, pt, bold), 1e-6)))
                cur = cur[cut:]
                lines += 1
    return lines


def autofit_pt(text: str, box_w_in: float, box_h_in: float, start_pt: float,
               min_pt: float, bold: bool = False, step: float = 0.5,
               max_lines: int | None = None) -> float:
    """Largest size in [min_pt, start_pt] at which `text` fits the box."""
    pt = start_pt
    while pt > min_pt:
        lines = wrap_count(text, pt, box_w_in, bold)
        h = lines * pt * LINE_FACTOR / 72.0
        if h <= box_h_in and (max_lines is None or lines <= max_lines):
            return pt
        pt = round(pt - step, 2)
    return min_pt


# --------------------------------------------------------------------------
# Fit log
# --------------------------------------------------------------------------

@dataclass
class FitEntry:
    slide: int
    label: str
    text: str
    pt: float
    bold: bool
    w_in: float
    h_in: float
    max_lines: int | None = None
    need_h: float | None = None   # precomputed (stacked paragraphs, tables)
    lines: int | None = None

    def evaluate(self) -> tuple[bool, str]:
        if self.need_h is not None:
            lines, need_h = (self.lines or 0), self.need_h
            shape = f"{lines} line(s), " if lines else ""
            if need_h > self.h_in + 1e-4:
                return False, (f"needs {need_h:.2f} in, box is {self.h_in:.2f} in "
                               f"({shape}stacked at {self.pt:g} pt)")
            return True, f"{shape}{need_h:.2f}/{self.h_in:.2f} in"
        lines = wrap_count(self.text, self.pt, self.w_in, self.bold)
        need_h = lines * self.pt * LINE_FACTOR / 72.0
        if self.max_lines is not None and lines > self.max_lines:
            return False, (f"{lines} lines > {self.max_lines} allowed "
                           f"(box {self.w_in:.2f} in at {self.pt:g} pt)")
        if need_h > self.h_in + 1e-4:
            return False, (f"needs {need_h:.2f} in, box is {self.h_in:.2f} in "
                           f"({lines} lines at {self.pt:g} pt)")
        return True, f"{lines} line(s), {need_h:.2f}/{self.h_in:.2f} in"


@dataclass
class FitLog:
    entries: list[FitEntry] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def add(self, slide: int, label: str, text: str, pt: float, w_in: float,
            h_in: float, bold: bool = False, max_lines: int | None = None) -> None:
        if text and text.strip():
            self.entries.append(FitEntry(slide, label, text, pt, bold, w_in, h_in, max_lines))

    def note(self, msg: str) -> None:
        self.notes.append(msg)

    def problems(self) -> list[tuple[FitEntry, str]]:
        out = []
        for e in self.entries:
            ok, why = e.evaluate()
            if not ok:
                out.append((e, why))
        return out

    def report(self, verbose: bool = False) -> str:
        lines = [f"Fit report — {metric_font_line()}"]
        bad = self.problems()
        if verbose:
            for e in self.entries:
                ok, why = e.evaluate()
                lines.append(f"  {'ok  ' if ok else 'OVER'} s{e.slide:>2} {e.label:<34} {why}")
        if bad:
            lines.append("")
            lines.append(f"{len(bad)} box(es) would overflow:")
            for e, why in bad:
                lines.append(f"  s{e.slide:>2} {e.label:<34} {why}")
                lines.append(f"        {e.text[:110]!r}")
        else:
            lines.append(f"  {len(self.entries)} text boxes checked — no overflow.")
        for n in self.notes:
            lines.append(f"  note: {n}")
        return "\n".join(lines)


# --------------------------------------------------------------------------
# Presentation helpers
# --------------------------------------------------------------------------

def open_base(base_path: str | Path) -> Presentation:
    prs = Presentation(str(base_path))
    strip_slides(prs)
    return prs


def strip_slides(prs: Presentation) -> None:
    """Remove every slide, keeping masters, layouts, theme and media."""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        try:
            prs.part.drop_rel(rId)
        except KeyError:
            pass
        sldIdLst.remove(sldId)


def pick_layout(prs: Presentation, *preferred: str):
    """First layout whose name matches one of `preferred` (case-insensitive),
    else the first layout that has a title placeholder, else layout 0."""
    layouts = [l for m in prs.slide_masters for l in m.slide_layouts]
    lower = {l.name.lower(): l for l in layouts}
    for name in preferred:
        if name.lower() in lower:
            return lower[name.lower()]
    for name in preferred:
        for l in layouts:
            if name.lower() in l.name.lower():
                return l
    for l in layouts:
        if any(ph.placeholder_format.idx == 0 for ph in l.placeholders):
            return l
    return layouts[0]


def new_slide(prs: Presentation, layout, title: str | None = None,
              header: str | None = None, keep_body: bool = False,
              title_pt: float | None = None, title_min_pt: float = 18.0,
              title_box: tuple[float, float, float, float] | None = None):
    slide = prs.slides.add_slide(layout)
    # Clear layout-inherited content placeholders we do not use.
    for ph in list(slide.placeholders):
        idx = ph.placeholder_format.idx
        if idx == 0:
            if title is None:
                ph._element.getparent().remove(ph._element)
            else:
                if title_box:
                    bx, by, bw, bh = title_box
                    ph.left, ph.top = inch(bx), inch(by)
                    ph.width, ph.height = inch(bw), inch(bh)
                pt = title_pt
                if pt is not None:
                    w = (ph.width or inch(12.05)) / EMU_IN
                    h = (ph.height or inch(0.43)) / EMU_IN
                    # Masters in this family uppercase the title; measure the
                    # uppercase form or the estimate under-reads by ~15%.
                    pt = autofit_pt(title.upper(), w, max(h, 0.44), pt,
                                    title_min_pt, bold=False, max_lines=1)
                set_text(ph, title, pt)
        elif idx == 4:      # slide number — keep, PowerPoint expects it
            continue
        elif idx == 34:     # running header
            if header is None:
                ph._element.getparent().remove(ph._element)
            else:
                set_text(ph, header)
        elif not keep_body:
            ph._element.getparent().remove(ph._element)
    return slide


def copy_slide_number(slide, layout) -> bool:
    """Put the layout's slide-number placeholder on the slide.

    python-pptx does not clone sldNum/ftr/dt placeholders when a slide is
    added, so a slide built this way carries no page number even though the
    layout defines one. The section-slide references all carry it; the sales
    deck's content slides do not — so this is opt-in per builder. The copy gets
    a fresh shape id, or PowerPoint reports the file as needing repair.
    """
    from copy import deepcopy
    from pptx.oxml.ns import qn
    for ph in layout.placeholders:
        if ph.placeholder_format.idx != 4:
            continue
        if any(p.placeholder_format.idx == 4 for p in slide.placeholders):
            return True
        el = deepcopy(ph._element)
        cNvPr = el.find(".//" + qn("p:cNvPr"))
        if cNvPr is not None:
            cNvPr.set("id", str(slide.shapes._next_shape_id))
        slide.shapes._spTree.append(el)
        return True
    return False


def set_text(placeholder, text: str, pt: float | None = None) -> None:
    tf = placeholder.text_frame
    tf.word_wrap = True
    tf.text = text
    if pt is not None:
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(pt)


# --------------------------------------------------------------------------
# Shapes
# --------------------------------------------------------------------------

def rect(slide, x, y, w, h, fill: str | None = None, line: str | None = None,
         line_pt: float = 0.75, rounded: bool = False, adj: float | None = None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        inch(x), inch(y), inch(w), inch(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = RGBColor.from_string(line)
        shp.line.width = Pt(line_pt)
    else:
        shp.line.fill.background()
    if rounded and adj is not None:
        shp.adjustments[0] = adj
    shp.shadow.inherit = False
    if shp.has_text_frame:
        shp.text_frame.word_wrap = True
    return shp


def panel(slide, x, y, w, h, fill=C["panel_grey"], accent: str | None = None,
          accent_w: float = 0.05, line: str | None = C["hairline"], rounded=False):
    """A tinted card with an optional left accent bar — the house idiom."""
    body = rect(slide, x, y, w, h, fill=fill, line=line, rounded=rounded)
    if accent:
        rect(slide, x, y, accent_w, h, fill=accent, line=None)
    return body


def line_h(slide, x, y, w, color=C["hairline"], thickness=0.01):
    return rect(slide, x, y, w, thickness, fill=color, line=None)


_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT,
          "j": PP_ALIGN.JUSTIFY}


def textbox(slide, x, y, w, h, paras: Sequence[dict], anchor: str = "t",
            wrap: bool = True):
    """Paragraph model:
        {"t": "text" | [("part", {...}), ...], "sz": 10, "b": False,
         "color": "26282B", "align": "l", "space_after": 0, "space_before": 0,
         "font": FONT_BODY, "spacing": 1.0}
    """
    tb = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN.get(spec.get("align", "l"), PP_ALIGN.LEFT)
        if spec.get("space_after"):
            p.space_after = Pt(spec["space_after"])
        if spec.get("space_before"):
            p.space_before = Pt(spec["space_before"])
        if spec.get("spacing"):
            p.line_spacing = spec["spacing"]
        parts = spec["t"]
        if isinstance(parts, str):
            parts = [(parts, {})]
        for text, over in parts:
            r = p.add_run()
            r.text = text
            style = {**spec, **over}
            r.font.size = Pt(style.get("sz", 10))
            r.font.bold = bool(style.get("b", False))
            r.font.name = style.get("font", FONT_BODY)
            col = style.get("color", C["ink"])
            if col:
                r.font.color.rgb = RGBColor.from_string(col)
    return tb


def para_text(spec: dict) -> str:
    parts = spec["t"]
    if isinstance(parts, str):
        return parts
    return "".join(t for t, _ in parts)


def stacked_height(paras: Sequence[dict], w: float,
                   default_sz: float = 10.0) -> tuple[float, int, float]:
    """(height in inches, total lines, largest point size) for a paragraph stack."""
    total, lines_total, biggest = 0.0, 0, default_sz
    for spec in paras:
        txt = para_text(spec)
        sz = spec.get("sz", default_sz)
        bold = bool(spec.get("b", False))
        n = wrap_count(txt, sz, w, bold)
        total += n * sz * LINE_FACTOR * float(spec.get("spacing", 1.0)) / 72.0
        total += (float(spec.get("space_after", 0))
                  + float(spec.get("space_before", 0))) / 72.0
        lines_total += n
        biggest = max(biggest, sz)
    return total, lines_total, biggest


def autofit_paras(paras: Sequence[dict], w: float, h: float,
                  min_scale: float = 0.72, step: float = 0.04,
                  default_sz: float = 10.0, max_scale: float = 1.0,
                  fill: float = 0.82) -> list[dict]:
    """Scale a paragraph stack proportionally to fit `h`.

    Scales DOWN until it fits. With `max_scale` > 1 it also scales UP until the
    stack fills `fill` of the box — slide-design rule 1: a box more than half
    empty means the type is too small, not that the box should stay airy.
    """
    def scaled_at(scale):
        return [{**p, "sz": round(p.get("sz", default_sz) * scale, 2)} for p in paras]

    scale = 1.0
    while True:
        scaled = scaled_at(scale)
        total, _, _ = stacked_height(scaled, w, default_sz * scale)
        if total <= h or scale <= min_scale:
            break
        scale = round(scale - step, 3)
    if max_scale > 1.0 and total < h * fill:
        while scale < max_scale:
            nxt = round(scale + step, 3)
            cand = scaled_at(nxt)
            total2, _, _ = stacked_height(cand, w, default_sz * nxt)
            if total2 > h * 0.96:
                break
            scale, scaled = nxt, cand
    return scaled


def log_box(fit: FitLog, slide_no: int, label: str, x, y, w, h,
            paras: Sequence[dict], default_sz: float = 10.0) -> None:
    """Register a multi-paragraph box with the fit log as one stacked estimate."""
    total, lines, biggest = stacked_height(paras, w, default_sz)
    text = " / ".join(para_text(p) for p in paras)[:200]
    fit.entries.append(FitEntry(slide_no, label, text, biggest, False, w, h,
                                need_h=total, lines=lines))


# --------------------------------------------------------------------------
# Pack-spec access
# --------------------------------------------------------------------------


# --------------------------------------------------------------------------
# The shared product catalog: ids in the spec, canonical names in the artifact
# --------------------------------------------------------------------------
# naming-and-clearance.md §1: "The catalog carries the canonical display name;
# artifacts render that, never a local variant." The spec carries ids, so a
# builder that prints `p["id"]` ships `oci-dedicated-ai-cluster` to a seller.

_CATALOG_CACHE: dict[str, dict[str, str]] | None = None


def _catalog_paths() -> list[Path]:
    here = Path(__file__).resolve()
    out = []
    for up in (3, 5):                       # plugin root, then a source checkout
        if len(here.parents) > up:
            out.append(here.parents[up] / "shared" / "data" / "oracle-products.yaml")
    env = os.environ.get("ORACLE_PRODUCT_CATALOG")
    if env:
        out.insert(0, Path(env))
    return out


def catalog() -> dict[str, dict[str, str]]:
    """{id: {"name": ..., "short": ...}} — empty when no catalog is reachable."""
    global _CATALOG_CACHE
    if _CATALOG_CACHE is None:
        _CATALOG_CACHE = {}
        for path in _catalog_paths():
            if not path.is_file():
                continue
            try:
                import yaml
                data = yaml.safe_load(path.read_text(encoding="utf-8"))
            except Exception:                # a broken catalog is not a crash
                continue
            rows = data.get("products") if isinstance(data, dict) else data
            for row in rows or []:
                if isinstance(row, dict) and row.get("id"):
                    _CATALOG_CACHE[str(row["id"])] = {
                        "name": str(row.get("name") or "").strip(),
                        "short": str(row.get("short") or "").strip()}
            break
    return _CATALOG_CACHE


def product_name(entry, short: bool = False) -> str:
    """Display name for an oracle_products[] entry or a bare catalog id.

    The spec's own `name` wins (a pack may have a reason), then the catalog,
    then the id — which at least resolves, and which the fit report shows.
    """
    if isinstance(entry, dict):
        local = str(entry.get("name") or "").strip()
        pid = str(entry.get("id") or "").strip()
    else:
        local, pid = "", str(entry or "").strip()
    if local:
        return local
    row = catalog().get(pid)
    if row:
        return (row["short"] if short and row["short"] else row["name"]) or pid
    return pid


class SpecError(RuntimeError):
    pass


class Spec:
    """Dotted read-only access to a pack spec, with channel-aware helpers."""

    def __init__(self, data: dict, channel: str = "partner_print"):
        self.data = data or {}
        self.channel = channel

    @classmethod
    def load(cls, path: str | Path, channel: str = "partner_print") -> "Spec":
        import yaml
        with open(path, "r", encoding="utf-8") as fh:
            data = yaml.safe_load(fh)
        if not isinstance(data, dict):
            raise SpecError(f"{path}: expected a YAML mapping at the top level")
        return cls(data, channel)

    def get(self, dotted: str, default=None):
        cur: Any = self.data
        for key in dotted.split("."):
            if isinstance(cur, dict) and key in cur:
                cur = cur[key]
            else:
                return default
            if cur is None:
                return default
        return cur

    def need(self, dotted: str):
        val = self.get(dotted)
        if val in (None, "", [], {}):
            raise SpecError(
                f"pack spec is missing `{dotted}` — run /oracle-packs:spec and "
                f"confirm that component before building this artifact")
        return val

    # ---- naming -------------------------------------------------------
    def name(self) -> str:
        variants = self.get("meta.name_variants", {}) or {}
        if self.channel == "internal":
            return variants.get("internal_slide") or variants.get("site") or self.need("meta.name")
        return variants.get("external") or variants.get("site") or self.need("meta.name")

    def subheading(self) -> str | None:
        if self.channel == "internal":
            return None
        return self.get("meta.name_variants.external_subheading")

    # ---- clearance ----------------------------------------------------
    def customer_name_allowed(self) -> bool:
        return bool(self.get(f"clearance.customer_name_allowed.{self.channel}", False))

    def customer_label(self) -> str:
        """What this channel may call the source customer."""
        if self.customer_name_allowed():
            named = self.get("meta.source_engagement.customer")
            if named:
                return str(named)
        return str(self.get("clearance.anonymized_descriptor",
                            "an enterprise customer"))

    def contact(self) -> dict:
        return (self.get(f"contacts.{self.channel}")
                or self.get("contacts.partner_print")
                or self.get("contacts.internal") or {})

    # ---- tiers --------------------------------------------------------
    def tiers(self) -> list[dict]:
        tiers = self.get("packages.tiers", []) or []
        if not tiers:
            raise SpecError("pack spec has no `packages.tiers` — the service "
                            "packages component is not signed off yet")
        return tiers

    def tier_by_id(self, tid: str) -> dict:
        for t in self.tiers():
            if t.get("id") == tid:
                return t
        return {}

    def tier_label(self, tier: dict, with_size: bool = True) -> str:
        name = tier.get("name") or tier.get("id", "").title()
        size = tier.get("size_tag")
        return f"{name}  ·  {size}" if (with_size and size) else name

    # ---- KPIs ---------------------------------------------------------
    def kpis(self) -> list[dict]:
        return self.get("kpis", []) or []

    # The spec writes `-` for "measured per engagement, no cleared number"
    # (shared/schema/pack-spec.md). Such a metric has no figure to print: a `-`
    # tile is an empty container reading as content, which rule 3 forbids.
    EMPTY = {"-", "--", "\u2014", "\u2013", "n/a"}

    @staticmethod
    def has_text(value) -> bool:
        """Filled, and not the spec's `-` marker for 'deliberately empty'."""
        text = str(value or "").strip()
        return bool(text) and text.lower() not in Spec.EMPTY

    @classmethod
    def has_figure(cls, kpi: dict) -> bool:
        fig = str((kpi or {}).get("figure") or "").strip()
        return bool(fig) and fig.lower() not in cls.EMPTY

    def figured_kpis(self) -> list[dict]:
        return [k for k in self.sales_kpis() if self.has_figure(k)]

    # ---- metric kinds --------------------------------------------------
    # `kind` decides where a metric may print (shared/schema/pack-spec.md):
    # business on the sales tiles and chips, leading only beside its business
    # metric, technical never — a proof-of-value acceptance criterion belongs to
    # the PoV package's success line. Absent means business, so a brief written
    # before the key existed keeps every tile it had (2026-09-23).
    @staticmethod
    def kpi_kind(kpi: dict) -> str:
        value = str((kpi or {}).get("kind") or "").strip().lower()
        return value if value in ("business", "leading", "technical") else "business"

    def sales_kpis(self) -> list[dict]:
        """The metrics a sales artifact may print: everything but the technical ones."""
        return [k for k in self.kpis() if self.kpi_kind(k) != "technical"]

    def technical_kpis(self) -> list[dict]:
        return [k for k in self.kpis() if self.kpi_kind(k) == "technical"]

    def pov_success_line(self) -> str:
        """"Proof accepted when: …" — the technical criteria, where the brief has any."""
        bits = []
        for k in self.technical_kpis():
            name = str(k.get("name") or "").strip()
            figure = str(k.get("figure") or "").strip()
            if not name:
                continue
            bits.append(f"{name} {figure}" if self.has_text(figure) else name)
        return ("Proof accepted when: " + ", ".join(bits)) if bits else ""

    def proof_word(self) -> str:
        """`Proven` only where a delivered result backs it (naming-and-clearance §3)."""
        delivered = any(k.get("figure_status") == "delivered_result" and self.has_figure(k)
                        for k in self.kpis())
        return "Proven" if delivered else "Proof of value"

    def kpi_attribution(self) -> str:
        for k in self.sales_kpis():
            att = k.get("attribution") or {}
            if self.customer_name_allowed() and att.get("named_when_allowed"):
                return str(att["named_when_allowed"])
            if att.get("otherwise"):
                return str(att["otherwise"])
        return f"proof of value at {self.customer_label()}"

    def kpi_caveat(self) -> str:
        for k in self.sales_kpis():
            if k.get("caveat"):
                return str(k["caveat"])
        return "Figures are illustrative and subject to confirmation."


# --------------------------------------------------------------------------
# Value formatting
# --------------------------------------------------------------------------

_CUR = {"EUR": "€", "USD": "$", "GBP": "£"}


def _money(v: float, cur: str) -> str:
    sym = _CUR.get(cur, (cur + " ") if cur else "")
    if v >= 1000 and v % 1000 == 0:
        return f"{sym}{int(v / 1000)}K"
    if v >= 1_000_000:
        return f"{sym}{v / 1_000_000:g}M"
    return f"{sym}{int(v):,}"


def fmt_price(price: dict | None, tbd: str = "To be defined") -> tuple[str, bool]:
    """Returns (rendered, needs_footnote)."""
    if not price:
        return tbd, False
    status = (price.get("status") or "").lower()
    cur = price.get("currency", "EUR")
    footnote = status in ("indicative", "estimate", "modeled")
    if price.get("range"):
        lo, hi = price["range"]
        lo_s = _money(lo, cur)
        hi_s = _money(hi, cur).lstrip(_CUR.get(cur, ""))
        return f"{lo_s}–{hi_s}", footnote
    if price.get("value") is not None:
        return _money(price["value"], cur), footnote
    return tbd, False


def fmt_duration(d: dict | None, tbd: str = "To be defined") -> str:
    if not d:
        return tbd
    lo, hi, target = d.get("min"), d.get("max"), d.get("target")
    if lo and hi and lo != hi:
        return f"{lo}–{hi} weeks"
    v = target or lo or hi
    return f"{v} weeks" if v else tbd


def wrap_to(text: str, width: int) -> str:
    import textwrap
    return "\n".join(textwrap.wrap(text, width))


def sibling_kit_path() -> Path:
    """Where a sibling skill's deckkit lives (for the import fallback)."""
    return Path(__file__).resolve().parents[2] / "deck" / "tools"
