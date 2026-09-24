#!/usr/bin/env python3
"""Build the 10-slide accelerator-pack sales deck by FILLING the exemplar deck.

    build_deck_v2.py <pack-spec.md> --out <dir> [--channel partner_print|internal]
                     [--fit-report] [--allow-overflow] [--exemplar <pptx>]

Stage 2 of docs/DECK-FIDELITY.md — "clone, don't redraw". The reference
Workforce Optimization deck ships as `assets/exemplar/wfo-sales-deck.pptx`; this
builder opens a copy of it in memory, keeps the slides the anatomy maps to,
duplicates the one slide type the reference lacks, and then only ever
**replaces content**: text runs keep their own formatting, pictures are swapped
inside their frames, table rows are cloned or removed. Geometry, fonts, corner
radii, colours, icons and the running header come from the exemplar by
construction — there is no drawing code here.

The slot map (semantic name -> shape id per slide) is
`assets/exemplar/slots.json`; the reasoning behind each mapping and the few
places where the builder does compute a number are in
`references/exemplar-builder.md`.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Sequence

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
for _up in range(2, 6):                     # the plugin's synced shared/tools, or the bundle's
    _shared = HERE.parents[_up] / "shared" / "tools" if len(HERE.parents) > _up else None
    if _shared is not None and (_shared / "build_diagram.py").is_file():
        sys.path.insert(0, str(_shared))
        break

from deckkit import (  # noqa: E402  (deliberate: import, never edit)
    FitLog, Spec, SpecError, fmt_duration, fmt_price, product_name,
)
import exemplar as ex  # noqa: E402
import build_diagram as diagram  # noqa: E402  (the one architecture model, shared by all three artifacts)

for _up in range(2, 6):                     # the spec stamp: the plugin's synced shared/tools, or the bundle's
    _shared = HERE.parents[_up] / "shared" / "tools" if len(HERE.parents) > _up else None
    if _shared is not None and (_shared / "spec_stamp.py").is_file():
        sys.path.insert(0, str(_shared))
        break
import spec_stamp  # noqa: E402  (which spec the deck was built from, written into its properties)

EXEMPLAR_DEFAULT = HERE.parent / "assets" / "exemplar" / "wfo-sales-deck.pptx"
SLOTS_DEFAULT = HERE.parent / "assets" / "exemplar" / "slots.json"
def _library_icon_map() -> Path:
    """The shared icon library's map, shared/data/icons/map.yaml, beside the plugin's tools."""
    for up in range(2, 6):
        if len(HERE.parents) > up:
            candidate = HERE.parents[up] / "shared" / "data" / "icons" / "map.yaml"
            if candidate.is_file():
                return candidate
    return HERE.parent / "assets" / "icons" / "map.yaml"     # absent: every card keeps its icon


ICON_MAP_DEFAULT = _library_icon_map()

HEADER_BRAND = "Oracle AI & Data Solutions"
# The retired lockup, in every casing it has been written in, anywhere in the header —
# the owner found it as "OCI accelerators" mid-string, not only as a prefix (2026-09-23).
LEGACY_HEADER_RE = re.compile(r"OCI\s+AI\s+Accelerators?|OCI\s+accelerators?", re.IGNORECASE)

TIER_ORDER = ("pov", "integration", "scaling")
TIER_DEFAULT_GLYPH = {"pov": "◐", "integration": "●", "scaling": "●●"}
GLYPH_KIND = {"◐": "partial", "●": "included", "●●": "advanced"}
EMPTY_MARKS = {"", "-", "--", "—", "–", "n/a"}


# --------------------------------------------------------------------------
# Small text utilities (content only — nothing here touches formatting)
# --------------------------------------------------------------------------

def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def sentences(text: str) -> list[str]:
    text = clean(text)
    if not text:
        return []
    parts = re.split(r"(?<=[.!?])\s+", text)
    return [p for p in parts if p]


def split_lead(text: str, seps: Sequence[str] = (":", " — ", " - ", ".")) -> tuple[str, str]:
    """Bold lead clause + regular remainder, the way every exemplar card is written."""
    text = clean(text)
    for sep in seps:
        if sep in text:
            head, tail = text.split(sep, 1)
            if 8 <= len(head) <= 110 and tail.strip():
                keep = head + (sep if sep in (":", ".") else sep.rstrip())
                return keep.strip(), (" " + tail.strip() if sep in (":", ".") else " " + tail.strip())
    return text, ""


def first_clause(text: Any, limit: int = 66) -> str:
    """The opening clause of a sentence — for captions that must stay one line."""
    out = clean(text)
    for sep in (";", " — ", ":", ". "):
        if sep in out:
            out = out.split(sep, 1)[0]
    out = out.rstrip(" .;:,")
    if len(out) > limit:
        out = out[:limit].rsplit(" ", 1)[0]
    return out


def is_empty(value: Any) -> bool:
    return clean(value).lower() in EMPTY_MARKS


# --------------------------------------------------------------------------
# Fit logging against the exemplar's own boxes
# --------------------------------------------------------------------------

HUGE = 99.0   # the height check is replaced by the line budget below


def _line_budget(pt: float, h_in: float) -> int:
    return max(1, int((h_in * 72.0) / (pt * 1.22) + 1e-6))


def log_shape(fit: FitLog, slide_no: int, label: str, shp, text: str,
              para: int = 0, run: int = 0, original: str | None = None,
              pt: float | None = None, w_in: float | None = None) -> None:
    """Record a filled slot against the exemplar's own box, type size and line count.

    The budget is the larger of (a) the lines the box holds and (b) the lines the
    exemplar's own text used in that slot — the exemplar is proof that N lines
    fit even where PowerPoint auto-sized the box to its content.

    `w_in` is the usable width to measure against, for the one case where the
    shape's own frame is not what renders: a shape inside a group carries its
    coordinates in the group's child space, and the group scales them.
    """
    if not clean(text):
        return
    pt = pt or ex.run_pt(shp, para, run) or 11.0
    bold = ex.run_bold(shp, para, run)
    w, h = ex.inner_box_in(shp)
    if w_in is not None:
        w = max(w_in, 0.1)
    allowed = _line_budget(pt, h)
    if original:
        from deckkit import wrap_count
        allowed = max(allowed, wrap_count(original, pt, w, bold))
    fit.add(slide_no, label, text, pt, w, HUGE, bold=bold, max_lines=allowed)


def log_cell(fit: FitLog, slide_no: int, label: str, cell, text: str,
             col_w: float, row_h: float, original: str | None = None,
             pt: float | None = None) -> None:
    if not clean(text):
        return
    pt = pt or ex.cell_run_pt(cell, 0, 0) or 10.0
    w = max(col_w - ex.to_in(cell.margin_left) - ex.to_in(cell.margin_right), 0.2)
    allowed = _line_budget(pt, max(row_h, 0.2))
    if original:
        from deckkit import wrap_count
        allowed = max(allowed, wrap_count(original, pt, w, False))
    fit.add(slide_no, label, text, pt, w, HUGE, max_lines=allowed)


# --------------------------------------------------------------------------
# Icons
# --------------------------------------------------------------------------

def load_icon_map(path: Path) -> list[tuple[list[str], Path]]:
    """Tolerant reader for assets/icons/map.yaml (owned by the stage-1 builder).

    Accepts either ``{file: [keywords]}``, ``{keyword: file}`` or a list of
    ``{file:, keywords: []}`` rows. Returns [(keywords, absolute path)].
    """
    if not path.is_file():
        return []
    try:
        import yaml
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except Exception:
        return []
    rows: list[tuple[list[str], Path]] = []

    def add(file: Any, keywords: Any) -> None:
        if not file:
            return
        p = Path(str(file))
        if not p.is_absolute():
            p = path.parent / p
        kws = keywords if isinstance(keywords, (list, tuple)) else [keywords]
        rows.append(([clean(k).lower() for k in kws if clean(k)], p))

    if isinstance(data, dict):
        data = data.get("icons", data)
    if isinstance(data, dict):
        for key, val in data.items():
            if isinstance(val, (list, tuple)):
                add(key, val)
            elif isinstance(val, dict):
                add(val.get("file") or val.get("path") or key,
                    val.get("keywords") or val.get("match") or [key])
            else:
                add(val, [key])
    elif isinstance(data, list):
        for row in data:
            if isinstance(row, dict):
                add(row.get("file") or row.get("path"),
                    row.get("keywords") or row.get("match") or [row.get("name")])
    return [(k, p) for k, p in rows if p.is_file()]


def pick_icon(icons, name: str) -> Path | None:
    text = clean(name).lower()
    best, score = None, 0
    for keywords, path in icons:
        hit = sum(len(k) for k in keywords if k and k in text)
        if hit > score:
            best, score = path, hit
    return best


# --------------------------------------------------------------------------
# Slide builders — each one fills an exemplar slide in place
# --------------------------------------------------------------------------

SPEC_PATH_HINT = None


def _image_path(value, spec_path=""):
    """A picture entry is a path string or the pictures step's record {file, source, ...}.

    A relative file is resolved against the brief's own folder when the caller knows
    where the brief is, and against the working directory otherwise.
    """
    if isinstance(value, dict):
        value = value.get("file")
    value = clean(value)
    if not value:
        return ""
    p = Path(value).expanduser()
    if not p.is_absolute() and spec_path:
        p = Path(spec_path).resolve().parent / p
    return str(p)


class Build:
    def __init__(self, spec: Spec, prs, slots: dict, fit: FitLog, icons):
        self.spec = spec
        self.spec_path = getattr(spec, "path", "") or ""   # set by main(); "" in library use
        self.prs = prs
        self.slots = slots
        self.fit = fit
        self.icons = icons
        self.notes: list[str] = []
        self.diagram: list[str] = []
        self._diagram_model: dict | None = None
        self.orig: dict[tuple[int, int], list[str]] = {}
        self.pictures: list[str] = []      # slots the owner still has to choose

    def note(self, msg: str) -> None:
        if msg in self.notes:                      # one line per finding, not per card
            return
        self.notes.append(msg)
        self.fit.note(msg)

    def note_technical_routing(self) -> None:
        """Say out loud that technical criteria were kept off the sales tiles."""
        technical = self.spec.technical_kpis()
        if not technical:
            return
        names = ", ".join(clean(k.get("name")) for k in technical if clean(k.get("name")))
        msg = (f"metrics: {len(technical)} technical criterion(s) ({names}) are kept off the "
               f"KPI chips and the proof tiles — they print on the service-packages slides as "
               f"the PoV's \"Proof accepted when …\" line.")
        if msg not in self.notes:
            print("build_deck_v2: " + msg, file=sys.stderr)
        self.note(msg)

    # -- fit logging against the exemplar's own text ----------------------
    def snapshot(self, slides) -> None:
        """Remember every slot's original text, before a single fill."""
        for i, slide in enumerate(slides, 1):
            for shp in ex.iter_shapes(slide):
                if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
                    self.orig[(i, shp.shape_id)] = [p.text for p in shp.text_frame.paragraphs]

    def log(self, slide_no: int, label: str, shp, text: str,
            para: int = 0, run: int = 0, pt: float | None = None,
            w_in: float | None = None) -> None:
        paras = self.orig.get((slide_no, shp.shape_id))
        original = paras[para] if paras and para < len(paras) else None
        log_shape(self.fit, slide_no, label, shp, text, para, run, original, pt, w_in)

    def logc(self, slide_no: int, label: str, cell, text: str, col_w: float,
             row_h: float, original: str | None = None, pt: float | None = None) -> None:
        log_cell(self.fit, slide_no, label, cell, text, col_w, row_h, original, pt)

    # -- shared -----------------------------------------------------------
    def stack(self) -> list[dict]:
        return list(self.spec.get("architecture.stack", []) or [])

    def layer_like(self, *words: str, default_index: int | None = None) -> dict:
        for layer in self.stack():
            hay = clean(layer.get("layer")).lower()
            if any(w in hay for w in words):
                return layer
        stack = self.stack()
        if default_index is not None and stack:
            return stack[max(-len(stack), min(default_index, len(stack) - 1))]
        return {}

    def named(self, key: str, fallback: str = "") -> str:
        """A spec string with `{customer}` / `{Customer}` filled with what this channel
        may call the source customer (the name when cleared, the descriptor otherwise)."""
        text = clean(self.spec.get(key)) or fallback
        who = self.spec.customer_label()
        return (text.replace("{Customer}", who[0].upper() + who[1:] if who else "")
                    .replace("{customer}", who))

    # -- 1 cover ----------------------------------------------------------
    def cover(self, slide) -> None:
        s = self.slots["slides"]["cover"]["slots"]
        title = ex.by_id(slide, s["cover.title"]["id"])
        name = self.spec.name()
        # The cover's own budget: the exemplar's one-liner is ~100 characters over
        # three lines at 25 pt. The short variant is what fits that block.
        one_liner = clean(self.spec.get("one_liner.short") or self.spec.get("one_liner.full"))
        ex.set_paragraphs(title, [(0, [name]), (1, [one_liner])])
        self.log(1, "cover.title", title, name, para=0)
        self.log(1, "cover.one_liner", title, one_liner, para=1)

        tier_shape = ex.find_id(slide, s["cover.tier_line"])
        if tier_shape is not None:
            icp = clean(self.spec.get("icp.line"))
            if icp:
                ex.fill_text(tier_shape, icp)
                self.log(1, "cover.icp", tier_shape, icp)
            else:
                ex.fill_text(tier_shape, "")
            self.note("cover: the reference's tier ladder line is removed "
                      "(no tier eyebrow); the shape now carries the ICP line.")
        sub = self.spec.subheading()
        if sub:
            self.note(f"cover: the exemplar cover has no subheading slot — "
                      f"`{sub}` is not printed on slide 1.")
        self.cover_photo(slide)

    def cover_photo(self, slide) -> None:
        """The cover hero is shared across the pack family — kept unless this pack chooses its own.

        It lives on the `Title-AI` layout, not on the slide. With
        `deck.images.cover` it is swapped in place (centre-cropped to the frame);
        without it, the family's own hero stays — the owner's rule (2026-09-23):
        "the title slide image can be shared across the packs". An ink-only
        cover is never the default.
        """
        layout = slide.slide_layout
        hero, area = None, 0.0
        for shp in layout.shapes:
            if shp.shape_type != 13 or "logo" in (shp.name or "").lower():
                continue
            size = ex.to_in(shp.width) * ex.to_in(shp.height)
            if size > area:
                hero, area = shp, size
        if hero is None:
            return
        path = _image_path((self.spec.get("deck.images", {}) or {}).get("cover"), self.spec_path)
        if path and Path(path).is_file():
            ex.replace_picture(hero, path)
            self.note(f"cover: hero image replaced from `deck.images.cover` ({path}).")
            return
        self.note("cover: the family hero kept — shared across the packs; set `deck.images.cover` only to replace it.")

    # -- 2 use case -------------------------------------------------------
    def use_case(self, slide) -> None:
        s = self.slots["slides"]["use_case"]["slots"]
        ex.fill_text(ex.by_id(slide, s["title"]), "USE CASE")

        body = ex.by_id(slide, s["use_case.problem.body"])
        protos = ex.paragraph_prototypes(body)
        lead, tail = split_lead(self.spec.need("problem_solution.problem"))
        raw_points = self.spec.get("problem_solution.problem_points") or []
        points = [f"{x['label']}: {x['text']}" if isinstance(x, dict) else x for x in raw_points]  # the schema writes {label, text}; the exemplar splits on the first ":" — normalize BEFORE clean(), which would stringify the dict
        points = [clean(p) for p in points if clean(p)]
        specs: list[tuple[int, list[str]]] = [(0, [lead, tail] if tail else [lead])]
        if points:
            specs.append((1, [""]))
            specs.append((2, ["This leads to the following challenges:"]))
            for pt in points:
                if ":" in pt:
                    head, rest = pt.split(":", 1)
                    specs.append((3, [head, ": " + rest.strip()]))
                else:
                    specs.append((3, [pt]))
        ex.set_paragraphs(body, specs, protos)
        self.log(2, "use_case.problem.body", body,
                  " ".join([lead, tail] + points), para=0)

        sol = ex.by_id(slide, s["use_case.solution.body"])
        sol_protos = ex.paragraph_prototypes(sol)
        reframe = clean(self.spec.get("problem_solution.reframe"))
        sents = sentences(self.spec.need("problem_solution.solution"))
        lead_txt = (reframe.rstrip(".") + ". ") if reframe else (sents.pop(0) + " " if sents else "")
        first = sents[0] if sents else ""
        rest = " ".join(sents[1:])
        sol_specs = [(0, [lead_txt, first] if first else [lead_txt])]
        if rest:
            sol_specs.append((1, [""]))
            sol_specs.append((2, [rest]))
        ex.set_paragraphs(sol, sol_specs, sol_protos)
        self.log(2, "use_case.solution.body", sol,
                  " ".join([lead_txt, first, rest]), para=0)

        chips = s["use_case.kpi_chips"]
        # Business metrics only: a chip is a claim about the buyer's business, and a
        # proof criterion ("reviewer agreement") reads as one when it sits in the row.
        labels = [clean(k.get("chip") or k.get("name")) for k in self.spec.sales_kpis()]
        labels = [l for l in labels if l]
        self.note_technical_routing()
        for i, chip in enumerate(chips):
            text_shape = ex.find_id(slide, chip["text"])
            if i < len(labels):
                ex.fill_text(text_shape, labels[i])
                self.log(2, f"use_case.kpi_chip[{i}]", text_shape, labels[i])
            else:
                ex.delete_ids(slide, [chip["text"], chip["box"]])
        if len(labels) > len(chips):
            self.note(f"use case: {len(labels)} KPI chips in the spec, {len(chips)} slots "
                      f"on the exemplar — kept the first {len(chips)}.")

        anchor = ex.by_id(slide, s["use_case.anchor_strip"])
        anchor_text = clean(self.spec.get("packages.anchor_line")) or self.anchor_fallback()
        a_lead, a_rest = split_lead(anchor_text, (" — ", " - ", ":"))
        ex.set_paragraphs(anchor, [(0, [a_lead, a_rest] if a_rest else [a_lead])])
        self.log(2, "use_case.anchor_strip", anchor, anchor_text)

    def anchor_fallback(self) -> str:
        required = [product_name(p) for p in (self.spec.get("oracle_products") or [])
                    if (p or {}).get("role") == "required"]
        if required:
            return "Anchored to " + " & ".join(required[:2]) + " — every run is OCI consumption."
        return "Anchored to the customer's Oracle estate — every run is OCI consumption."

    # -- 3 verticals ------------------------------------------------------
    def verticals(self, slide) -> None:
        s = self.slots["slides"]["verticals"]["slots"]
        ex.fill_text(ex.by_id(slide, s["title"]), "VERTICAL APPLICATIONS")
        cards = s["verticals"]
        rows = list(self.spec.get("verticals", []) or [])
        from pptx.dml.color import RGBColor
        for i, card in enumerate(cards):
            name_shape = ex.find_id(slide, card["name"])
            body_shape = ex.find_id(slide, card["body"])
            icon_shape = ex.find_id(slide, card["icon"])
            if i >= len(rows):
                ex.fill_text(name_shape, "")
                ex.fill_text(body_shape, "")
                ex.delete_ids(slide, [card["rule"]])
                if icon_shape is not None:
                    ex.delete_shape(icon_shape)
                panel = ex.find_id(slide, card["panel"])
                if panel is not None:
                    panel.fill.solid()
                    panel.fill.fore_color.rgb = RGBColor.from_string("DCE1E5")
                self.note(f"verticals: slot {i + 1} has no vertical — greyed empty card "
                          f"(rule 3), never a gap.")
                continue
            row = rows[i]
            name = clean(row.get("name"))
            body = clean(row.get("what_matters_here") or
                         (row.get("framing") or {}).get("solution"))
            ex.fill_text(name_shape, name)
            ex.fill_text(body_shape, body)
            self.log(3, f"verticals[{i}].name", name_shape, name)
            self.log(3, f"verticals[{i}].body", body_shape, body)
            if icon_shape is None:
                continue
            icon = self.vertical_icon(row, name)
            if icon:
                ex.replace_picture(icon_shape, icon, cover=False)
            elif self.icons:
                self.note(f"verticals: no icon chosen for '{name}' and no library keyword "
                          f"matched — kept the exemplar's icon in that slot.")
            else:
                self.note("verticals: no icon chosen and the icon library is not present — "
                          "every card keeps the exemplar's own icon.")
        if len(rows) > len(cards):
            self.note(f"verticals: {len(rows)} in the spec, {len(cards)} cards on the "
                      f"exemplar — kept the first {len(cards)}.")

    def vertical_icon(self, row: dict, name: str) -> Path | None:
        """The owner's pick from the pictures step (`verticals[].icon`), in the white render
        the reference deck's cards carry; else the shared library's keyword match."""
        chosen = row.get("icon")
        if chosen:
            if isinstance(chosen, dict):
                chosen = chosen.get("file_white") or chosen.get("file")
            path = Path(_image_path(chosen, self.spec_path))
            if path.is_file():
                return path
            self.note(f"verticals: the icon chosen for '{name}' is not at {path} — "
                      f"used the library's instead.")
        return pick_icon(self.icons, name)

    # -- 4 today -> tomorrow ----------------------------------------------
    def today_tomorrow(self, slide) -> None:
        s = self.slots["slides"]["today_tomorrow"]["slots"]
        images = self.spec.get("deck.images", {}) or {}
        logo = ex.find_id(slide, s["today_tomorrow.customer_logo"])
        headline_shape = ex.by_id(slide, s["today_tomorrow.headline"])
        self.place_logo(slide, logo, headline_shape, images, slide_no=4)

        headline = clean(self.spec.get("problem_solution.reframe_question") or
                         self.spec.get("problem_solution.reframe"))
        ex.fill_text(headline_shape, headline)
        self.log(4, "today_tomorrow.headline", headline_shape, headline)

        sub = ex.by_id(slide, s["today_tomorrow.subline"])
        sub_text = clean(self.spec.get("one_liner.short") or self.spec.get("one_liner.full"))
        ex.fill_text(sub, sub_text)
        self.log(4, "today_tomorrow.subline", sub, sub_text)

        case = ex.by_id(slide, s["today_tomorrow.vertical_case"])
        verticals = self.spec.get("verticals", []) or []
        case_text = self.named("deck.vertical_case") or (
            f"Vertical case: {clean(verticals[0].get('name'))}" if verticals else "")
        ex.fill_text(case, case_text)
        self.log(4, "today_tomorrow.vertical_case", case, case_text)

        today = clean(self.spec.get("problem_solution.today") or
                      self.spec.get("problem_solution.problem"))
        tomorrow = clean(self.spec.get("problem_solution.tomorrow") or
                         self.spec.get("problem_solution.solution"))
        for key, text, label in (("today.text", today, "today.text"),
                                 ("tomorrow.text", tomorrow, "tomorrow.text")):
            shp = ex.by_id(slide, s[key])
            ex.fill_text(shp, text)
            self.log(4, label, shp, text)

        for key, spec_key, caption in (
                ("today.image", "today", "Before — image to be chosen"),
                ("tomorrow.image", "tomorrow", "After — image to be chosen")):
            pic = ex.find_id(slide, s[key])
            if pic is None:
                continue
            path = _image_path(images.get(spec_key), self.spec_path)
            if path and Path(path).is_file():
                ex.replace_picture(pic, path)
            else:
                ex.empty_picture_frame(slide, pic, caption,
                                       caption_proto=ex.find_id(slide, s["today.text"]))
                self.pictures.append(
                    f"{spec_key:<10} s4  deck.images.{spec_key} — the frame is an explicit "
                    f"empty container until the real before/after screen is dropped in")

    def place_logo(self, slide, logo, headline_shape, images: dict, slide_no: int) -> None:
        """The source customer's logo: only when cleared AND supplied."""
        if logo is None:
            return
        path = _image_path(images.get("customer_logo"), self.spec_path)
        if self.spec.customer_name_allowed() and path and Path(path).is_file():
            ex.replace_picture(logo, path, cover=False)
            return
        ex.delete_shape(logo)
        x, y, w, h = ex.frame_in(headline_shape)
        right = x + w
        ex.set_frame(headline_shape, x=0.42, w=max(right - 0.42, 1.0))
        self.note(f"s{slide_no}: the reference's customer logo is removed "
                  f"(clearance) and the headline moved to the left margin.")

    # -- 5 proof ----------------------------------------------------------
    def proof(self, slide) -> None:
        s = self.slots["slides"]["proof"]["slots"]
        images = self.spec.get("deck.images", {}) or {}
        headline_shape = ex.by_id(slide, s["proof.headline"])
        self.place_logo(slide, ex.find_id(slide, s["proof.customer_logo"]),
                        headline_shape, images, slide_no=5)

        headline = self.named("deck.proof_headline") or clean(
            self.spec.get("one_liner.short") or self.spec.get("one_liner.full"))
        ex.fill_text(headline_shape, headline)
        self.log(5, "proof.headline", headline_shape, headline)

        self.fill_stats(slide, s["proof.stats"])

        blocks = s["proof.blocks"]
        # The exemplar's four blocks, in its own words: the customer's situation, what
        # was built for them, and the value to each side (WfO slide 5, kept verbatim
        # at the owner's request — 2026-09-23).
        customer = self.spec.customer_label()
        context = self.named("meta.source_engagement.context") or (customer[0].upper() + customer[1:])
        solution = self.named("meta.source_engagement.delivered")
        for_partner = (self.named("packages.value_for_partner")
                       or clean(self.spec.get("packages.anchor_line")) or self.anchor_fallback())
        for_client = (self.named("packages.value_for_client")
                      or clean(self.spec.get("problem_solution.solution")))
        content = [
            ("CONTEXT", context),
            ("SOLUTION", solution),
            ("VALUE FOR ORACLE + NVIDIA", for_partner),
            ("VALUE FOR CLIENT", for_client),
        ]
        for i, block in enumerate(blocks):
            label, text = content[i]
            lab = ex.by_id(slide, block["label"])
            bod = ex.by_id(slide, block["body"])
            ex.fill_text(lab, label)
            ex.fill_text(bod, text)
            self.log(5, f"proof.block[{i}].label", lab, label)
            self.log(5, f"proof.block[{i}].body", bod, text)

        anchor = ex.by_id(slide, s["proof.anchor_strip"])
        anchor_text = clean(self.spec.get("packages.anchor_line")) or self.anchor_fallback()
        a_lead, a_rest = split_lead(anchor_text, (" — ", " - ", ":"))
        ex.set_paragraphs(anchor, [(0, [a_lead, a_rest] if a_rest else [a_lead])])
        self.log(5, "proof.anchor_strip", anchor, anchor_text)

        foot = ex.by_id(slide, s["proof.footnote"])
        if self.spec.figured_kpis():
            foot_text = f"Figures from {self.spec.kpi_attribution()}. {self.spec.kpi_caveat()}"
        else:   # nothing to attribute yet: say where the first engagement stands instead
            foot_text = f"First engagement: {self.spec.kpi_attribution()} — contracted, results to follow."
        ex.fill_text(foot, clean(foot_text))
        self.log(5, "proof.footnote", foot, foot_text)

    def fill_stats(self, slide, stats) -> None:
        self.note_technical_routing()
        figured = self.spec.figured_kpis()
        # Business metrics only; with no figure yet the strip names what the proof of
        # value measures — still business metrics, never its acceptance criteria.
        kpis = figured or self.spec.sales_kpis()
        restricted = [k for k in self.spec.sales_kpis()
                      if k.get("channels") and self.spec.channel not in (k.get("channels") or [])]
        if restricted:
            for stat in stats:
                ex.delete_ids(slide, [stat["box"], stat["value"], stat["label"]])
            self.note("proof: a metric is not cleared for this channel — peer claims are "
                      "all or none (rule 11), so the whole stat strip is dropped.")
            return
        filled: list[tuple] = []
        for i, stat in enumerate(stats):
            if i >= len(kpis):
                ex.delete_ids(slide, [stat["box"], stat["value"], stat["label"]])
                continue
            kpi = kpis[i]
            if self.spec.has_figure(kpi):
                figure = clean(kpi.get("figure"))
                if kpi.get("show_baseline") and clean(kpi.get("baseline")):
                    figure = f"{clean(kpi['baseline'])} → {figure}"
                label = clean(kpi.get("label") or kpi.get("formula") or kpi.get("name"))
            else:   # the metric's name where the figure will stand, and what it measures
                figure = re.sub(r"\s*[↑↓→]+\s*$", "", clean(kpi.get("chip"))) or clean(kpi.get("name"))
                label = (clean(kpi.get("label") or kpi.get("name"))
                         + " — to be measured in the proof of value")   # the same words as the one-pager and the executive summary
            v = ex.by_id(slide, stat["value"])
            l = ex.by_id(slide, stat["label"])
            ex.fill_text(v, figure)
            ex.fill_text(l, label)
            filled.append((i, v, l, figure, label))
        # A metric name is longer than a figure: if any value needs a smaller size, every
        # value takes it — the three boxes are peers (slide-design rule 2).
        from pptx.util import Pt
        pt = None
        for step in (None, 16.0, 14.0):
            if step is not None:
                for _, v, _, _, _ in filled:
                    for p in v.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(step)
                pt = step
            if all(self.text_height(v, [fig], ex.to_in(v.width)) <= ex.to_in(v.height) + 0.01
                   for _, v, _, fig, _ in filled):
                break
        for i, v, l, figure, label in filled:
            self.log(5, f"proof.stat[{i}].value", v, figure, pt=pt)
            self.log(5, f"proof.stat[{i}].label", l, label)
        if not figured:
            self.note("proof: no cleared figure — the stat strip names the metrics to be measured.")

    # -- 6 why it sells (duplicated from the proof slide) -------------------
    def why_it_sells(self, slide) -> None:
        s = self.slots["slides"]["proof"]["slots"]
        ex.delete_ids(slide, [s["proof.customer_logo"]])
        headline_shape = ex.by_id(slide, s["proof.headline"])
        headline = clean(self.spec.get("deck.seller_lead")) or (
            "What this pack gives an account exec that a custom project does not.")
        x, y, w, h = ex.frame_in(headline_shape)
        ex.set_frame(headline_shape, x=0.42, w=max(x + w - 0.42, 1.0))
        ex.fill_text(headline_shape, headline)
        self.log(6, "why.headline", headline_shape, headline)

        # Top strip: the tier ladder — what a seller can actually put in front of an
        # account. Names and scope only: timings and prices live on the package
        # slides, where their footnote travels with them.
        for i, stat in enumerate(s["proof.stats"]):
            tiers = self.spec.tiers()
            if i >= len(tiers):
                ex.delete_ids(slide, [stat["box"], stat["value"], stat["label"]])
                continue
            tier = tiers[i]
            value = self.spec.tier_label(tier)
            caption = first_clause(tier.get("scope_line") or
                                   (tier.get("what_you_get") or [""])[0])
            v = ex.by_id(slide, stat["value"])
            l = ex.by_id(slide, stat["label"])
            ex.fill_text(v, value)
            ex.fill_text(l, caption)
            self.log(6, f"why.tier[{i}].name", v, value)
            self.log(6, f"why.tier[{i}].caption", l, caption)

        claims = [clean(c) for c in (self.spec.get("packages.why_it_sells_for_the_partner") or [])
                  if clean(c)]
        raw_consumption = self.spec.get("packages.target_oci_consumption")
        consumption = clean(raw_consumption) if self.spec.has_text(raw_consumption) else (
            "To be defined" if raw_consumption is not None else "")   # `-` = deliberately empty: keep the panel, say so; an absent key drops it
        blocks = s["proof.blocks"]
        filled: list[tuple[str, str]] = []
        for claim in claims[:len(blocks) - (1 if consumption else 0)]:
            lead, rest = split_lead(claim, (" — ", " - ", ":"))
            label = lead.rstrip(" —-:·").upper()
            body = rest.strip() or claim
            filled.append((label, body[0].upper() + body[1:] if body else body))
        if consumption and len(filled) < len(blocks):
            filled.append(("OCI CONSUMPTION", consumption))
        for slot_ix, block_ix in enumerate(range(len(blocks))):
            block = blocks[block_ix]
            lab = ex.by_id(slide, block["label"])
            bod = ex.by_id(slide, block["body"])
            if slot_ix >= len(filled):
                ex.delete_ids(slide, [block["bar"], block["label"],
                                      block["panel"], block["body"]])
                continue
            label, body = filled[slot_ix]
            ex.fill_text(lab, label)
            ex.fill_text(bod, body)
            self.log(6, f"why.card[{slot_ix}].label", lab, label)
            self.log(6, f"why.card[{slot_ix}].body", bod, body)
        if len(claims) > len(blocks):
            self.note(f"why it sells: {len(claims)} claims in the spec, {len(blocks)} cards "
                      f"on the exemplar — kept the first {len(filled)}.")

        cta = ex.by_id(slide, s["proof.anchor_strip"])
        contact = self.spec.contact()
        who = " · ".join(x for x in (clean(contact.get("name")), clean(contact.get("title")),
                                     clean(contact.get("email") or contact.get("mailbox"))) if x)
        question = clean(self.spec.get("deck.cta")) or "Ready to test the fit in one of your accounts?"
        ex.set_paragraphs(cta, [(0, [question, "   " + who] if who else [question])])
        self.log(6, "why.cta", cta, question + " " + who)

        foot = ex.by_id(slide, s["proof.footnote"])
        if self.spec.figured_kpis():   # the caveat travels with the figures; without them the line has nothing to say
            ex.fill_text(foot, "Figures indicative and subject to confirmation; tiers, timing and "
                               "prices on the service-packages slides.")
        else:
            ex.delete_shape(foot)

    # -- 7 solution layers -------------------------------------------------
    def layers(self, slide) -> None:
        cfg = self.slots["slides"]["layers"]
        s = cfg["slots"]
        ex.fill_text(ex.by_id(slide, s["title"]), "TECHNOLOGY STACK")   # the reference's own title
        sub = ex.by_id(slide, s["layers.subline"])
        sub_text = clean(self.spec.get("deck.layers_sub")) or (
            f"How {self.spec.name()} is layered — from the infrastructure it runs on "
            f"to the customer's own rules.")
        ex.fill_text(sub, sub_text)
        self.log(7, "layers.subline", sub, sub_text)

        rows_cfg = s["layers.rows"]
        stack = self.stack()
        if not stack:
            raise SpecError("pack spec has no `architecture.stack` — the solution-layers "
                            "slide has nothing to show")
        # The ladder is the technology stack from the business application down
        # (the reference: application · engine · infrastructure). A layer the spec
        # places ABOVE the application — the client's own configuration — is the
        # application row's "Tailored solution" card, never a row of its own.
        app = self.layer_like("app", "business", "by softserve", default_index=0)
        app_ix = next((i for i, l in enumerate(stack) if l is app), 0)
        for layer in stack[:app_ix]:
            self.note(f"solution layers: `{clean(layer.get('layer'))}` sits above the "
                      f"application layer — it is the Tailored solution card, not a row "
                      f"(the reference's ladder).")
        stack = stack[app_ix:]
        # The exemplar's ladder band and gap; row heights are distributed inside it.
        top = ex.to_in(ex.by_id(slide, rows_cfg[0]["panel"]).top)
        last = ex.by_id(slide, rows_cfg[-1]["panel"])
        bottom = ex.to_in(last.top) + ex.to_in(last.height)
        gap = 0.16
        want = len(stack)
        proto_cfg = rows_cfg[1]

        # Grow the ladder by cloning the plain row prototype (its panel + children).
        # Clones go in *above* the last row, so infrastructure stays at the bottom of
        # the ladder and the exemplar's vendor tints keep their order.
        rows: list[dict] = [dict(r) for r in rows_cfg]
        while len(rows) < want:
            clone: dict = {}
            for key in ("panel", "name", "divider", "summary", "badge"):
                sid = proto_cfg.get(key)
                if not sid:
                    continue
                src = ex.by_id(slide, sid)
                clone[key] = ex.clone_shape(slide, src).shape_id
            rows.insert(len(rows) - 1, clone)
        while len(rows) > want:
            dead = rows.pop(max(len(rows) - 2, 0))
            ex.delete_ids(slide, [v for v in dead.values() if isinstance(v, int)])
            for card in dead.get("cards", []) or []:
                ex.delete_ids(slide, [v for v in card.values() if isinstance(v, int)])
            ex.delete_ids(slide, dead.get("chips", []) or [])

        row_h = (bottom - top - gap * (want - 1)) / want
        for i, (row, layer) in enumerate(zip(rows, stack)):
            panel = ex.by_id(slide, row["panel"])
            old_top, old_h = ex.to_in(panel.top), ex.to_in(panel.height)
            new_top = top + i * (row_h + gap)
            k = (row_h / old_h) if old_h else 1.0
            members = [row.get(k_) for k_ in ("name", "divider", "summary", "badge")]
            for card in row.get("cards", []) or []:
                members.extend(card.values())
            members.extend(row.get("chips", []) or [])
            for sid in members:
                if not isinstance(sid, int):
                    continue
                shp = ex.find_id(slide, sid)
                if shp is None:
                    continue
                # A resized row resizes what it holds: offsets scale with the row,
                # boxes, chips and dividers scale their height too; a text box keeps
                # its height and auto-fits (2026-09-23: cards left at the reference's
                # height poked out of a shorter band).
                off = ex.to_in(shp.top) - old_top
                shp.top = ex.inch(new_top + off * k)
                if abs(k - 1.0) > 0.005 and getattr(shp, "shape_type", None) != 17:
                    shp.height = ex.inch(ex.to_in(shp.height) * k)
            panel.top = ex.inch(new_top)
            panel.height = ex.inch(row_h)

            name_shape = ex.by_id(slide, row["name"])
            ex.fill_lines(name_shape, [clean(layer.get("layer")),
                                       clean(layer.get("vendor"))])
            self.log(7, f"layers[{i}].name", name_shape,
                      clean(layer.get("layer")), para=0)
            summary = clean(layer.get("summary")) or ", ".join(
                clean(x) for x in (layer.get("items") or []))
            if row.get("summary"):
                sm = ex.by_id(slide, row["summary"])
                ex.fill_text(sm, summary)
                self.log(7, f"layers[{i}].summary", sm, summary)
            if row.get("badge"):
                bd = ex.by_id(slide, row["badge"])
                ex.fill_text(bd, clean(layer.get("vendor")))
                self.log(7, f"layers[{i}].badge", bd, clean(layer.get("vendor")))
            for card in row.get("cards", []) or []:
                self.fill_layer_card(slide, card, layer, bottom=new_top + row_h)
            chips = row.get("chips") or []
            small_pt = None   # the caption size, taken from the first chip and shared by its peers
            for ci, chip_id in enumerate(chips):
                tiers = self.spec.tiers()
                chip = ex.find_id(slide, chip_id)
                if chip is None:
                    continue
                if ci >= len(tiers):
                    ex.delete_shape(chip)
                    continue
                tier = tiers[ci]
                # The reference's chip: the tier's name and two small lines of scope.
                scope = first_clause(tier.get("scope_line") or
                                     (tier.get("what_you_get") or [""])[0])
                if len(scope) > 40:   # two small lines: cut at a clause, never mid-phrase
                    scope = scope[:40].rsplit(", ", 1)[0] if ", " in scope[:40] else scope[:40].rsplit(" ", 1)[0]
                ex.fill_lines(chip, [clean(tier.get("name"))] + ([scope] if scope else []))
                self.log(7, f"layers.chip[{ci}]", chip, clean(tier.get("name")))
                paras = chip.text_frame.paragraphs
                if scope and len(paras) > 1:
                    # The reference's first chip holds its caption as a second, smaller
                    # paragraph; the other two hold it after a soft break in one
                    # paragraph, so their filled second paragraph would inherit the
                    # name's size. Peers share one caption size: the first chip's.
                    if small_pt is None:
                        p0, p1 = ex.last_run_pt(paras[0]._p), ex.last_run_pt(paras[1]._p)
                        small_pt = p1 if (p0 and p1 and p1 < p0) else 6.0
                    from pptx.util import Pt
                    for r in paras[1].runs:
                        r.font.size = Pt(small_pt)
                    self.log(7, f"layers.chip[{ci}].scope", chip, scope, para=1, pt=small_pt)

        arrow = ex.find_id(slide, s.get("layers.value_arrow"))
        if arrow is not None:
            arrow.top = ex.inch(top)
            arrow.height = ex.inch(bottom - top)
        if want != len(rows_cfg):
            self.note(f"solution layers: the spec has {want} layers, the exemplar ladder has "
                      f"{len(rows_cfg)} — rows cloned from the exemplar's own row and the band "
                      f"redistributed.")

    def fill_layer_card(self, slide, card: dict, layer: dict,
                        bottom: float | None = None) -> None:
        title = ex.find_id(slide, card.get("title"))
        body = ex.find_id(slide, card.get("body"))
        if title is not None and body is not None:
            # The reference's application row: the pack, pre-built and reusable, on
            # the left card; the tailored solution and its tiers on the right one.
            # The reference says "accelerator pack" / "packaged"; the print cut may not
            # (packaging vocabulary stays internal), so the card keeps the reference's
            # shape with the seller's words.
            ex.fill_lines(title, ["Oracle AI accelerator", "ORACLE / SOFTSERVE-ENHANCED"])
            summary = clean(layer.get("summary")) or ", ".join(
                clean(x) for x in (layer.get("items") or []) if clean(x))
            text = (f"Pre-built and reusable: {summary}" if summary
                    else f"Pre-built and reusable for {self.spec.name()} use cases")
            ex.fill_text(body, text)
            # The exemplar's body is one line that grows to fit; give it the room the
            # text needs inside its row, and step down to 8 pt only when 9 pt would
            # leave the row — so the fit report measures the real box.
            avail = ((bottom - ex.to_in(body.top) - 0.06) if bottom
                     else ex.to_in(body.height))
            pt = None
            need = self.text_height(body, [text], ex.to_in(body.width))
            if need > avail + 0.01:
                from pptx.util import Pt
                for p in body.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8)
                pt = 8.0
                need = self.text_height(body, [text], ex.to_in(body.width))
            body.height = ex.inch(min(max(ex.to_in(body.height), need), max(avail, 0.15)))
            self.log(7, "layers.card.pack", body, text, pt=pt)
        elif title is not None:
            ex.fill_lines(title, ["Tailored solution", "SOFTSERVE"])

    # -- 8 architecture ----------------------------------------------------
    # The exemplar draws two columns: the systems the pack reads on the left, the
    # platform container on the right.  A pack whose result lands in a system that
    # is not also a source needs a third column, so the container is narrowed from
    # the right and the destination boxes stand in the freed strip as peers of the
    # source boxes — same prototype, same width, same height (slide-design rule 2).
    # Both system columns are set to one width, and the corridor either side of the
    # container carries the arrow labels at the exemplar's own label width.
    LABEL_CORRIDOR = 1.30          # in, between a system column and the container
    PLATFORM_MIN_W = 4.75          # in, below this the app box stops holding its name
    PLATFORM_BOTTOM_MAX = 6.80     # in, the footer rule sits at 7.1
    COLUMN_MIN_W = 1.90            # in, below this a system name stops reading

    def architecture(self, slide) -> None:
        """Lay the pack's ONE architecture model into the exemplar's frame.

        Nothing here derives a node list: `shared/tools/build_diagram.py` did that
        once, for all three artifacts. This slide prints the model's `name` and
        `line` levels; the `detail` level belongs to the mini-site.
        """
        s = self.slots["slides"]["architecture"]["slots"]
        model = self.diagram_model()
        ex.fill_text(ex.by_id(slide, s["title"]), "ARCHITECTURE")
        sub = ex.by_id(slide, s["architecture.subline"])
        sub_text = clean(self.spec.get("deck.architecture_sub")) or (
            f"Reference architecture for a {self.spec.name()} implementation")
        gate = model.get("gate")
        if gate:
            # The human gate rides the subline as a small caption — the slide has no
            # footnote of its own, and the invariant is what the buyer asks about.
            tail = clean(f"{gate['name']} — {model.get('note') or gate.get('line', '')}")
            long_form = f"{sub_text} · {tail}"
            sub_text = long_form if len(long_form) <= 150 else f"{sub_text} · {gate['name']}"
        ex.fill_text(sub, sub_text)
        self.log(8, "architecture.subline", sub, sub_text)

        # What every box will say — decided before anything is placed, because the
        # platform's own geometry is laid out to what it has to hold.
        app = ex.by_id(slide, s["architecture.app"])
        app_name, app_sub = model["app"]["name"], model["app"]["line"]

        engine_text = ex.find_id(slide, s["architecture.engine.text"])
        engine_name = model["engine"]["name"]
        engine_sub = model["engine"]["line"]

        infra = ex.by_id(slide, s["architecture.infra"])
        infra_name = model["platform"]["label"]
        infra_sub = ", ".join(model["platform"]["services"])

        plan = self.flow_plan()
        geom = self.reflow_platform(slide, s, plan,
                                    [(app, [app_name, app_sub]),
                                     (engine_text, [engine_name, engine_sub]),
                                     (infra, [infra_name, infra_sub])])

        ex.fill_lines(app, [app_name, app_sub])
        self.log(8, "architecture.app", app, app_name, para=0)
        self.log(8, "architecture.app.sub", app, app_sub, para=1)
        if engine_text is not None:
            ex.fill_lines(engine_text, [engine_name, engine_sub])
            # The engine text lives inside a group, so its own frame is in the
            # group's child space — measure it against what actually renders.
            self.log(8, "architecture.engine", engine_text, engine_name, para=0,
                     w_in=geom["engine_w"])
            self.log(8, "architecture.engine.sub", engine_text, engine_sub, para=1,
                     w_in=geom["engine_w"])
        ex.fill_lines(infra, [infra_name, infra_sub])
        self.log(8, "architecture.infra", infra, infra_sub, para=1)

        ex.delete_ids(slide, [s["architecture.stray_rule"]])
        self.draw_flows(slide, s, plan, geom)

    def diagram_model(self) -> dict:
        """The pack's architecture model, built from the spec by the one function the
        one-pager and the mini-site also call, so all three draw the very same nodes
        for this spec and this channel. Never stored, so never stale.
        """
        if self._diagram_model is None:
            try:
                self._diagram_model = diagram.build_model(self.spec.data,
                                                          channel=self.spec.channel)
            except diagram.DiagramError as err:
                raise SpecError(f"the architecture picture cannot be drawn: {err}")
            for warning in self._diagram_model.get("warnings") or []:
                self.note(f"architecture: {warning}.")
        return self._diagram_model

    def flow_plan(self) -> dict:
        """The model's sources and destination-only systems, in the deck's own words.

        A system that is both an input and an output is ONE box with two arrows (a
        write-back, rule: never a reversed arrow); a system that only receives is a
        destination and gets its own box on the right.
        """
        model = self.diagram_model()
        sources = [{"system": row["name"], "in": row["data"], "out": ""}
                   for row in model["sources"]]
        dests: list[dict] = []
        for row in model["destinations"]:
            match = next((n for n in sources
                          if n["system"].lower() == row["name"].lower()), None)
            if match:
                match["out"] = row["data"]
            else:
                dests.append({"system": row["name"], "in": "", "out": row["data"]})
        if not sources and not dests:
            raise SpecError("pack spec has no `architecture.inputs` or `architecture.outputs` "
                            "— the architecture slide has no flows to draw")
        return {"sources": sources, "dests": dests}

    @staticmethod
    def text_height(shp, texts: Sequence[str], w_in: float, scale: float = 1.0) -> float:
        """Inches this shape needs to hold `texts` (one per paragraph) at `w_in`."""
        from deckkit import wrap_count
        tf = shp.text_frame
        pad_w = (ex.to_in(tf.margin_left) + ex.to_in(tf.margin_right)) * scale
        pad_h = (ex.to_in(tf.margin_top) + ex.to_in(tf.margin_bottom)) * scale
        inner = max(w_in - pad_w, 0.3)
        total = 0.0
        for i, text in enumerate(texts):
            if not clean(text):
                continue
            pt = ex.run_pt(shp, i, 0) or 11.0
            bold = ex.run_bold(shp, i, 0)
            total += wrap_count(clean(text), pt, inner, bold) * pt * 1.22 / 72.0
        return total + pad_h

    def reflow_platform(self, slide, s: dict, plan: dict, boxes) -> dict:
        """Place the columns and the platform container, then return their geometry.

        Two columns when every output goes back to a system the pack also reads
        (the container keeps the exemplar's full width); three when the spec names
        a destination-only system.
        """
        container = ex.by_id(slide, s["architecture.container"])
        node_proto = ex.by_id(slide, s["architecture.node_proto"])
        group = ex.find_id(slide, s["architecture.node_group"])
        engine = ex.find_id(slide, s["architecture.engine"])
        app, infra = boxes[0][0], boxes[2][0]

        c_x, c_y, c_w, c_h = ex.frame_in(container)
        n_x, n_y, n_w, n_h = ex.frame_in(node_proto)
        right_edge = c_x + c_w                      # the slide's own content margin
        band_top = ex.to_in(group.top) if group is not None else n_y
        col_w, dest_x = n_w, None

        if plan["dests"]:
            span = right_edge - n_x
            col_w = max(self.COLUMN_MIN_W,
                        min(n_w, (span - self.PLATFORM_MIN_W - 2 * self.LABEL_CORRIDOR) / 2))
            dest_x = right_edge - col_w
            new_x = n_x + col_w + self.LABEL_CORRIDOR
            new_w = dest_x - self.LABEL_CORRIDOR - new_x
            k = new_w / c_w
            inside = [app, engine, infra] + [
                ex.find_id(slide, sid) for sid in (s.get("architecture.engine_arrows") or [])]
            for shp in inside:
                if shp is None:
                    continue
                x, _, w, _ = ex.frame_in(shp)
                ex.set_frame(shp, x=new_x + (x - c_x) * k, w=w * k)
            ex.set_frame(container, x=new_x, w=new_w)
            ex.set_frame(node_proto, w=col_w)       # every box is cloned from it
            c_x, c_w = new_x, new_w

        # The platform's three blocks keep the exemplar's pads and gaps; a block
        # whose text needs more room grows, and the container grows with it.
        blocks = [(app, boxes[0][1]), (engine, boxes[1][1]), (infra, boxes[2][1])]
        frames = [ex.frame_in(shp) for shp, _ in blocks if shp is not None]
        pads = [frames[0][1] - c_y]
        for i in range(len(frames) - 1):
            pads.append(frames[i + 1][1] - (frames[i][1] + frames[i][3]))
        pads.append((c_y + c_h) - (frames[-1][1] + frames[-1][3]))

        eng_scale = ex.group_scale_x(engine) if engine is not None else 1.0
        widths = [frames[0][2],
                  ex.to_in(engine.width) if engine is not None else frames[1][2],
                  frames[2][2]]
        measured = [self.text_height(app, boxes[0][1], widths[0]),
                    self.text_height(boxes[1][0], boxes[1][1], widths[1], scale=eng_scale)
                    if boxes[1][0] is not None else 0.0,
                    self.text_height(infra, boxes[2][1], widths[2])]
        heights = [max(frames[i][3], measured[i]) for i in range(3)]
        total = sum(pads) + sum(heights)
        if total > c_h + 1e-6:
            grow = min(total - c_h, self.PLATFORM_BOTTOM_MAX - (c_y + c_h))
            c_h += max(grow, 0.0)
            if total > c_h + 1e-6:                  # the rest comes off the top pad
                pads[0] = max(0.25, pads[0] - (total - c_h))
                total = sum(pads) + sum(heights)
            if total > c_h + 1e-6:
                self.note("architecture: the platform boxes hold more text than the "
                          "container's band — shorten the app, engine or "
                          "infrastructure summary in the pack brief.")
            ex.set_frame(container, h=c_h)
        if plan["dests"] or any(heights[i] > frames[i][3] + 1e-6 for i in range(3)):
            y = c_y + pads[0]
            for i, (shp, _) in enumerate(blocks):
                if shp is not None:
                    ex.set_frame(shp, y=y, h=heights[i])
                y += heights[i] + pads[i + 1]
            self.restack_engine_arrows(slide, s, app, engine, c_x, c_w)

        band_bottom = max(n_y + n_h, c_y + c_h - 0.05)
        app_x, app_y, app_w, app_h = ex.frame_in(app)
        engine_w = widths[1]
        if engine is not None and boxes[1][0] is not None:
            tf = boxes[1][0].text_frame
            engine_w = ex.to_in(engine.width) - (
                ex.to_in(tf.margin_left) + ex.to_in(tf.margin_right)) * eng_scale
        return {"col_x": n_x, "col_w": col_w, "dest_x": dest_x,
                "band": (band_top, band_bottom),
                "container": (c_x, c_y, c_w, c_h),
                "app": (app_x, app_y, app_w, app_h),
                "engine_w": max(engine_w, 0.3)}

    def restack_engine_arrows(self, slide, s: dict, app, engine, c_x: float,
                              c_w: float) -> None:
        """The two short arrows between the app and the engine follow their boxes."""
        ids = s.get("architecture.engine_arrows") or []
        if engine is None or not ids:
            return
        _, app_y, _, app_h = ex.frame_in(app)
        top = app_y + app_h
        height = max(ex.to_in(engine.top) - top, 0.08)
        for sid in ids:
            shp = ex.find_id(slide, sid)
            if shp is None:
                continue
            ex.set_frame(shp, y=top, h=height)

    def draw_flows(self, slide, s: dict, plan: dict, geom: dict) -> None:
        """One box per system, one labelled arrow per data flow (rule 7)."""
        protos = {"node": ex.by_id(slide, s["architecture.node_proto"]),
                  "in": ex.by_id(slide, s["architecture.arrow_in"]),
                  "out": ex.by_id(slide, s["architecture.arrow_out"]),
                  "elbow": ex.by_id(slide, s["architecture.arrow_elbow"]),
                  "label": ex.by_id(slide, s["architecture.arrow_label_proto"])}
        lab_proto_w = ex.to_in(protos["label"].width)
        lab_h = ex.to_in(protos["label"].height)
        ex.delete_ids(slide, [s["architecture.node_group"], s["architecture.node_proto"],
                              s["architecture.arrow_in"], s["architecture.arrow_out"],
                              s["architecture.arrow_elbow"],
                              s["architecture.arrow_label_proto"]])
        for sid in (50, 62):                      # the reference's other two arrow labels
            ex.delete_ids(slide, [sid])

        col_x, col_w, dest_x = geom["col_x"], geom["col_w"], geom["dest_x"]
        band_top, band_bottom = geom["band"]
        c_x, _, c_w, _ = geom["container"]
        app_x, app_y, app_w, app_h = geom["app"]
        sources, dests = plan["sources"], plan["dests"]

        # One height for every system box, so a source and a destination are peers.
        v_gap, flow_sep = 0.25, 0.62
        tall = max(len(sources), len(dests), 1)
        node_h = (band_bottom - band_top - v_gap * (tall - 1)) / tall

        def column(n: int) -> list[float]:
            """The tops of `n` boxes, centred in the band."""
            span = n * node_h + (n - 1) * v_gap
            top = band_top + ((band_bottom - band_top) - span) / 2
            return [top + i * (node_h + v_gap) for i in range(n)]

        def place_box(node: dict, x: float, y: float) -> None:
            box = ex.clone_shape(slide, protos["node"], x, y, col_w, node_h)
            # The box carries the model's `name` and nothing else: what flows travels
            # on the arrow (rule 7), and a source's `detail` is the mini-site's level.
            name = node["system"]
            ex.fill_lines(box, [name])
            self.log(8, f"architecture.{node['role']}[{node['ix']}]", box, name, para=0)

        def label(x: float, y: float, w: float, text: str, tag: str) -> None:
            """The data on the arrow, sitting just above its own line."""
            w = max(w, 0.7)
            h = min(max(self.text_height(protos["label"], [text], w), 0.18), lab_h)
            lab = ex.clone_shape(slide, protos["label"], x, y - h - 0.06, w, h)
            ex.fill_text(lab, text)
            self.log(8, tag, lab, text)

        # -- the left column: every source, with its write-back where the spec has one
        gap_x1, gap_x2 = col_x + col_w, app_x
        lab_x = gap_x1 + 0.12
        bend_min = min(gap_x2 - 0.10, lab_x + lab_proto_w + 0.08)
        # A label stays in the corridor: it never runs over the platform container.
        lab_w = min(lab_proto_w, bend_min - lab_x - 0.08, c_x - lab_x - 0.08)
        adj0 = max(0.30, min(0.92, (bend_min - gap_x1) / max(gap_x2 - gap_x1, 0.1)))
        left_arrows = sum(1 for n in sources if n["in"]) + sum(1 for n in sources if n["out"])
        a_ix = 0
        for i, (node, y) in enumerate(zip(sources, column(len(sources)))):
            node.update({"role": "source", "ix": i})
            place_box(node, col_x, y)
            centre = y + node_h / 2
            flows = ([("in", node["in"])] if node["in"] else []) + \
                    ([("out", node["out"])] if node["out"] else [])
            for j, (direction, data) in enumerate(flows):
                a_ix += 1
                y_app = app_y + a_ix * app_h / (left_arrows + 1)
                # A write-back pair leaves the box as two lines, far enough apart
                # that each one's label sits clear above its own.
                y_node = centre + (j - (len(flows) - 1) / 2) * flow_sep
                cxn = ex.clone_shape(slide, protos["in" if direction == "in" else "out"])
                if abs(y_app - y_node) > 0.02:
                    ex.set_connector_geom(cxn, "bentConnector3",
                                          adj=int(min(adj0 * 100000 + (a_ix - 1) * 2000,
                                                      94000)))
                ex.place_connector(cxn, gap_x1, y_node, gap_x2, y_app)
                label(lab_x, y_node, lab_w, data, f"architecture.arrow.in[{a_ix}]")

        # -- the right column: one box per destination-only system
        if dests:
            app_right = app_x + app_w
            bend_x = c_x + c_w + 0.15
            run = max(dest_x - app_right, 0.2)
            lab_x_r = bend_x + 0.06
            lab_w_r = min(lab_proto_w, dest_x - lab_x_r - 0.06)
            for i, (node, y) in enumerate(zip(dests, column(len(dests)))):
                node.update({"role": "destination", "ix": i})
                place_box(node, dest_x, y)
                centre = y + node_h / 2
                y_app = app_y + (i + 1) * app_h / (len(dests) + 1)
                cxn = ex.clone_shape(slide, protos["in"])   # the head is at the far end
                if abs(centre - y_app) > 0.02:
                    ex.set_connector_geom(
                        cxn, "bentConnector3",
                        adj=int(max(0.06, min(0.90, (bend_x - app_right) / run))
                                * 100000 + i * 2000))
                ex.place_connector(cxn, app_right, y_app, dest_x, centre)
                label(lab_x_r, centre, lab_w_r, node["out"],
                      f"architecture.arrow.out[{i + 1}]")

        # The plain-sentence summary the review pack carries is the model's own:
        # the slide and the summary cannot disagree, because both are the model.
        self.diagram = diagram.describe(self.diagram_model()).splitlines()

    def refit_table(self, table, frame, col_w: list[float], slide_no: int,
                    footnote=None, grown: bool = False) -> None:
        """Keep a grown table inside the exemplar's band without shrinking any type.

        Adding capability rows makes the table taller than the frame the reference
        gave it. Rather than scale the type down (the finding this rewrite exists to
        fix), reclaim the slack the exemplar left in its own rows: a row whose
        declared height exceeds what its text needs gives the difference back,
        tallest slack first. If it still does not fit, that is reported, never hidden.
        """
        from deckkit import wrap_count
        rows = list(table.rows)
        band = ex.to_in(frame.height)
        need: list[float] = []
        for r, row in enumerate(rows):
            tallest = 0.35
            for c in range(len(col_w)):
                cell = table.cell(r, c)
                text = cell.text
                if not clean(text):
                    continue
                # A glyph cell is "<glyph 16 pt>  <prose 9 pt>": the prose run sets
                # the wrap, not the glyph.
                pt = ex.cell_run_pt(cell, 0, 0) or 10.0
                if len(cell.text_frame.paragraphs[0].runs) > 1:
                    pt = ex.last_run_pt(cell.text_frame.paragraphs[0]._p) or pt
                w = max(col_w[c] - ex.to_in(cell.margin_left) - ex.to_in(cell.margin_right), 0.2)
                lines = wrap_count(text, pt, w)
                h = lines * pt * 1.22 / 72.0 + ex.to_in(cell.margin_top) \
                    + ex.to_in(cell.margin_bottom)
                tallest = max(tallest, h)
            need.append(tallest)
        have = [ex.to_in(row.height) for row in rows]
        cushion = 1.04            # PowerPoint's own row box runs a little over the estimate
        total = sum(max(h, n) for h, n in zip(have, need)) * cushion
        over = total - band
        if over > 0:
            slack = sorted(((have[i] - need[i], i) for i in range(len(rows))
                            if have[i] - need[i] > 0.02), reverse=True)
            for amount, i in slack:
                if over <= 0:
                    break
                give = min(amount, over)
                rows[i].height = ex.inch(have[i] - give)
                total -= give
                over -= give
        # A taller table needs its legend out of the way. A table that grew past the
        # reference's own row count drops the legend to the slide's bottom line
        # outright — the renderer gives a row a little more than any estimate, and a
        # legend under the last row is not worth a guess.
        floor_y = 6.90
        bottom = ex.to_in(frame.top) + total
        if footnote is not None:
            want = floor_y if grown else min(bottom + 0.06, floor_y)
            if want > ex.to_in(footnote.top):
                footnote.top = ex.inch(want)
        if bottom + 0.06 > floor_y + 0.02:
            self.note(f"s{slide_no}: the table runs {bottom + 0.06 - floor_y:.2f} in past the "
                      f"reference band with every row already at its text height — cut a "
                      f"capability row or shorten the wording; nothing was scaled down.")

    # -- 9 / 10 package tables ---------------------------------------------
    def packages(self, slide, detailed: bool) -> None:
        key = "packages_detailed" if detailed else "packages"
        cfg = self.slots["slides"][key]
        s, tcfg = cfg["slots"], cfg["table"]
        slide_no = 10 if detailed else 9
        ex.fill_text(ex.by_id(slide, s["title"]),
                     "SERVICE PACKAGES (DETAILED)" if detailed else "SERVICE PACKAGES")
        frame = ex.by_id(slide, s[f"{key}.table"])
        table = frame.table
        col_w = [ex.to_in(c.width) for c in table.columns]
        tiers = self.spec.tiers()
        if len(tiers) != len(col_w) - 1:
            self.note(f"s{slide_no}: the exemplar table has {len(col_w) - 1} tier columns and "
                      f"the spec has {len(tiers)} tiers — filled the first "
                      f"{min(len(tiers), len(col_w) - 1)}; columns are not cloned.")

        rows = tcfg["rows"]
        was = [[table.cell(r, c).text for c in range(len(col_w))]
               for r in range(len(table.rows))]
        rows_before = len(table.rows)
        glyph_protos = {kind: table.cell(r, c)
                        for kind, (r, c) in tcfg["glyph_cells"].items()}
        glyph_protos = {k: ex.cell_paragraph_prototypes(v) for k, v in glyph_protos.items()}
        cap_proto_cells = [ex.cell_paragraph_prototypes(table.cell(tcfg["capability_proto"], c))
                           for c in range(len(col_w))]

        for ci, tier in enumerate(tiers[:len(col_w) - 1], start=1):
            label = self.spec.tier_label(tier)
            cell = table.cell(rows["header"], ci)
            ex.fill_cell(cell, label)
            self.logc(slide_no, f"packages.tier[{ci}]", cell, label,
                     col_w[ci], ex.to_in(table.rows[rows["header"]].height),
                     original=was[rows["header"]][ci])

        needs_footnote = False
        scope_row = rows["scope"]
        # Technical acceptance criteria never reach a sales tile; this is where they
        # belong — the PoV's own scope, as the line that says when the proof is done.
        success = self.spec.pov_success_line()
        for ci, tier in enumerate(tiers[:len(col_w) - 1], start=1):
            scope = clean(tier.get("scope_line")) or clean(
                (tier.get("what_you_get") or [""])[0])
            if success and str(tier.get("id") or "").strip() == "pov":
                scope = (scope.rstrip(". ") + ". " + success).strip() if scope else success
                self.note(f"s{slide_no}: the PoV column carries the proof's success line "
                          f"({len(self.spec.technical_kpis())} technical criterion(s)).")
            cell = table.cell(scope_row, ci)
            if detailed:
                ex.fill_cell(cell, scope)
            else:
                lead, rest = split_lead(scope, (".", ":", " — "))
                ex.set_cell(cell, [(0, [lead, rest] if rest else [lead])])
            self.logc(slide_no, f"packages.scope[{ci}]", cell, scope,
                     col_w[ci], ex.to_in(table.rows[scope_row].height),
                     original=was[scope_row][ci])

        if not detailed:
            for row_key, getter in (("price_services", "services_price"),
                                    ("price_infra", "infra_price_monthly")):
                for ci, tier in enumerate(tiers[:len(col_w) - 1], start=1):
                    text, foot = fmt_price(tier.get(getter))
                    if row_key == "price_infra" and text not in ("To be defined",):
                        text = f"{text} / month"
                    if foot:
                        text += "*"
                        needs_footnote = True
                    ex.fill_cell(table.cell(rows[row_key], ci), text)
            for ci, tier in enumerate(tiers[:len(col_w) - 1], start=1):
                ex.fill_cell(table.cell(rows["timing"], ci),
                             fmt_duration(tier.get("duration_weeks")))

        handling = list(self.spec.get("packages.capability_handling", []) or [])
        first = rows["capability_first"]
        ex.match_row_count(table, len(handling), tcfg["capability_proto"], first)
        for ri, entry in enumerate(handling):
            r = first + ri
            area = clean(entry.get("area"))
            label_cell = table.cell(r, tcfg["label_col"])
            ex.set_cell(label_cell, [(0, [area])], cap_proto_cells[tcfg["label_col"]])
            self.logc(slide_no, f"packages.area[{ri}]", label_cell, area,
                     col_w[0], ex.to_in(table.rows[r].height),
                     original=was[r][0] if r < len(was) else None)
            for ci, tier in enumerate(tiers[:len(col_w) - 1], start=1):
                tid = str(tier.get("id") or TIER_ORDER[min(ci - 1, len(TIER_ORDER) - 1)])
                prose = entry.get(tid, "")
                override = (entry.get("glyphs") or {}).get(tid)
                cell = table.cell(r, ci)
                if is_empty(prose) and not override:
                    ex.set_cell(cell, [(0, ["—"])], glyph_protos["none"])
                    continue
                lead = re.match(r"^(●●|●|◐|—)\s*", clean(prose))     # the spec may carry the glyph in the prose itself
                glyph = clean(override) or (lead.group(1) if lead else TIER_DEFAULT_GLYPH.get(tid, "●"))
                prose_text = clean(prose)[lead.end():] if lead else clean(prose)
                protos = (glyph_protos["none"] if glyph == "—" else
                          glyph_protos.get(GLYPH_KIND.get(glyph, "included"), glyph_protos["included"]))
                if detailed:
                    k = max(ex.distinct_runs(protos[0]), 2)
                    texts = (([glyph, "  "] + [""] * (k - 3) + [prose_text]) if k >= 3
                             else [glyph + "  ", prose_text])   # one spacer whatever the prototype's run count
                    ex.set_cell(cell, [(0, texts)], protos)
                    self.logc(slide_no, f"packages.cell[{ri}][{ci}]", cell,
                             prose_text, col_w[ci], ex.to_in(table.rows[r].height),
                             original=was[r][ci] if r < len(was) else None,
                             pt=ex.last_run_pt(protos[0]) or 9.0)
                else:
                    ex.set_cell(cell, [(0, [glyph])], protos)

        foot = ex.find_id(slide, s[f"{key}.footnote"])
        grown = len(table.rows) > rows_before
        self.refit_table(table, frame, col_w, slide_no, footnote=foot, grown=grown)
        if foot is not None:
            legend = (clean(self.spec.get("packages.capability_handling_legend")).replace(" · ", "     ")
                      or "◐ partial     ● included     ●● multi-region / advanced")   # the spec's own legend, in the exemplar's spacing
            if needs_footnote:
                # A grown table leaves room for one footnote line, not two.
                legend += ("   ·   " if grown else "\n")
                legend += ("* Indicative price; real consumption depends on run frequency, "
                           "rule-set complexity and the integration landscape.")
            ex.set_paragraphs(foot, [(i, [line]) for i, line in enumerate(legend.split("\n"))])


# --------------------------------------------------------------------------
# Assembly
# --------------------------------------------------------------------------

def running_header(spec: Spec) -> tuple[str, str | None]:
    template = clean(spec.get("deck.running_header"))
    note = None
    if template:
        m = LEGACY_HEADER_RE.search(template)
        if m:
            template = LEGACY_HEADER_RE.sub(HEADER_BRAND, template)
            note = (f"running header: the spec still carries the retired "
                    f"'{m.group(0)}' lockup — rewritten to '{HEADER_BRAND}'.")
    else:
        template = HEADER_BRAND + " — {name}"
    try:
        return template.format(name=spec.name()), note
    except Exception:
        return template, note


def build(spec: Spec, exemplar_path: Path, slots_path: Path, out_dir: Path,
          fit: FitLog, icon_map: Path, stamp: str | None = None) -> tuple[Path, Build]:
    slots = json.loads(Path(slots_path).read_text(encoding="utf-8"))
    prs = ex.open_exemplar(exemplar_path)
    icons = load_icon_map(icon_map)
    b = Build(spec, prs, slots, fit, icons)

    # The one slide type the reference lacks: duplicate the composition it derives
    # from (the proof slide's quadrant blocks), then order to the anatomy's ten.
    ex.duplicate_slide(prs, 4)                      # 0-based: exemplar slide 5
    order = [0, 1, 2, 3, 4, 10, 5, 6, 8, 9]         # cover..packages detailed; ref 8 dropped
    ex.arrange(prs, order)

    slides = list(prs.slides)
    b.snapshot(slides)
    header, header_note = running_header(spec)
    if header_note:
        b.note(header_note)
    header_id = slots["header"]["id"]
    for i, slide in enumerate(slides, 1):
        ex.clear_notes(slide)
        if i == 1:
            continue
        shp = ex.find_id(slide, header_id)
        if shp is not None:
            ex.fill_text(shp, header)

    steps = [
        ("cover", lambda s: b.cover(s)),
        ("use case", lambda s: b.use_case(s)),
        ("verticals", lambda s: b.verticals(s)),
        ("today -> tomorrow", lambda s: b.today_tomorrow(s)),
        ("proof", lambda s: b.proof(s)),
        ("why it sells", lambda s: b.why_it_sells(s)),
        ("solution layers", lambda s: b.layers(s)),
        ("architecture", lambda s: b.architecture(s)),
        ("service packages", lambda s: b.packages(s, detailed=False)),
        ("service packages (detailed)", lambda s: b.packages(s, detailed=True)),
    ]
    for i, (label, fn) in enumerate(steps):
        try:
            fn(slides[i])
        except SpecError:
            raise
        except Exception as exc:                    # keep a partial deck rather than nothing
            b.note(f"s{i + 1} ({label}) failed: {exc.__class__.__name__}: {exc}")

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{spec.get('meta.slug', 'pack')}-sales-deck.pptx"
    if stamp:
        prs.core_properties.identifier = stamp      # which spec this file reflects (CON006 / CON007)
    prs.save(str(out))
    return out, b


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="Build the accelerator-pack sales deck by filling the exemplar deck.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Done means: fit report clean, contact sheet compared against the exemplar, "
               "linter clean.")
    ap.add_argument("spec", help="path to pack-spec.md")
    ap.add_argument("--out", required=True, help="output directory")
    ap.add_argument("--channel", default="partner_print",
                    choices=["partner_print", "internal"],
                    help="who the deck is for; drives naming, clearance and contact")
    ap.add_argument("--exemplar", default=str(EXEMPLAR_DEFAULT),
                    help="the exemplar deck (default: the skill's asset)")
    ap.add_argument("--slots", default=str(SLOTS_DEFAULT), help="the slot map")
    ap.add_argument("--icons", default=str(ICON_MAP_DEFAULT), help="vertical icon map")
    ap.add_argument("--fit-report", action="store_true",
                    help="print the per-box text-fit estimate for every filled slot")
    ap.add_argument("--allow-overflow", action="store_true",
                    help="exit 0 even when boxes overflow (review builds only)")
    args = ap.parse_args(argv)

    try:
        spec = Spec.load(args.spec, channel=args.channel)
        spec.path = args.spec          # relative picture paths resolve against the brief's folder
    except SpecError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    fit = FitLog()
    built_from = spec_stamp.stamp(args.spec)
    try:
        out, b = build(spec, Path(args.exemplar), Path(args.slots), Path(args.out),
                       fit, Path(args.icons), stamp=built_from)
    except SpecError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2

    print(f"built {out}  (10 slides from the exemplar, channel={args.channel})")
    print(f"spec stamp: {built_from}")
    if b.diagram:
        print("architecture diagram:")
        for line in b.diagram:
            print(line)
    if b.pictures:
        print("pictures still to choose (the deck ships an explicit empty state for each):")
        for line in b.pictures:
            print("  " + line)
    print(fit.report(verbose=args.fit_report))
    if fit.problems() and not args.allow_overflow:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
