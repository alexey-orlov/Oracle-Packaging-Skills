#!/usr/bin/env python3
"""Checks for the exemplar builder — run it directly:

    python3 tests/test_build_deck_v2.py

Builds both test specs and asserts what "clone, don't redraw" means in the file:
ten slides in the anatomy's order, the running header rewritten to the current
lockup, no tier eyebrow, the family's cover hero still on the cover layout (it is
kept unless `deck.images.cover` replaces it — the owner's rule, 2026-09-23), the
exemplar's own type sizes and rounded corners still in place, and an architecture
rebuilt from the spec.
Stdlib + python-pptx only, no pytest required.
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "tools"))

from pptx import Presentation                       # noqa: E402
from pptx.oxml.ns import qn                         # noqa: E402

BUILDER = SKILL / "tools" / "build_deck_v2.py"
FIXTURE = HERE / "fixture-pack-spec.md"
VARIABILITY = HERE / "fixture-pack-spec_v2-variability.md"

failures: list[str] = []
checks = 0


def check(condition: bool, message: str) -> None:
    global checks
    checks += 1
    if not condition:
        failures.append(message)


def build(spec: Path, out: Path) -> tuple[int, str]:
    proc = subprocess.run(
        [sys.executable, str(BUILDER), str(spec), "--out", str(out)],
        capture_output=True, text=True)
    return proc.returncode, proc.stdout + proc.stderr


def all_shapes(container):
    for shp in container.shapes:
        yield shp
        if shp.shape_type == 6:
            yield from all_shapes(shp)


def texts(slide) -> str:
    out = []
    for shp in all_shapes(slide):
        if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            out.append(shp.text_frame.text)
        if getattr(shp, "has_table", False) and shp.has_table:
            for row in shp.table.rows:
                for cell in row.cells:
                    out.append(cell.text)
    return "\n".join(out)


def rounded(slide) -> int:
    return sum(1 for shp in all_shapes(slide)
               for geom in shp._element.iter(qn("a:prstGeom"))
               if geom.get("prst") == "roundRect")


def run_fixture(tmp: Path) -> None:
    out = tmp / "fixture"
    rc, log = build(FIXTURE, out)
    check(rc == 0, f"fixture build exited {rc}, not 0:\n{log}")
    check("no overflow" in log, f"fixture fit report is not clean:\n{log}")
    check("PoV" not in log.split("pictures")[0].split("architecture")[0],
          "unexpected tier vocabulary in the build banner")
    deck = out / "workforce-optimization-sales-deck.pptx"
    check(deck.is_file(), "fixture deck was not written")
    prs = Presentation(str(deck))
    slides = list(prs.slides)

    check(len(slides) == 10, f"{len(slides)} slides, expected 10")
    check(round(prs.slide_width / 914400, 2) == 13.33,
          "canvas is not the exemplar's 13.33 in")

    # The running header: the retired lockup never ships, and slide 1 has none.
    for i, slide in enumerate(slides[1:], 2):
        body = texts(slide)
        check("Oracle AI & Data Solutions — Workforce optimization" in body,
              f"slide {i} does not carry the current running header")
        check("OCI AI Accelerators" not in body,
              f"slide {i} still carries the retired OCI AI Accelerators lockup")

    cover = texts(slides[0])
    check("PROOF OF VALUE" not in cover.upper() or "ROLL-OUT" not in cover.upper(),
          "the cover still carries the reference's tier ladder line")
    check("Workforce optimization" in cover, "the cover does not name the pack")

    # The cover hero is shared across the pack family (the owner, 2026-09-23) — kept
    # unless the pack sets deck.images.cover, which the fixture does not.
    hero = [shp for shp in slides[0].slide_layout.shapes
            if shp.shape_type == 13 and "logo" not in (shp.name or "").lower()]
    check(bool(hero), "the family's cover hero is missing from the cover layout")

    # Fidelity: the exemplar's own rounded cards and type sizes are untouched.
    check(rounded(slides[1]) == 6,
          f"use-case slide has {rounded(slides[1])} rounded shapes, exemplar has 6")
    table = next(shp.table for shp in slides[8].shapes
                 if getattr(shp, "has_table", False) and shp.has_table)
    header = table.cell(0, 1).text_frame.paragraphs[0].runs[0]
    check(header.font.size.pt == 12.5,
          f"package table header is {header.font.size.pt} pt, exemplar is 12.5 pt")
    check("PoV Jumpstart" in table.cell(0, 1).text, "tier names not filled from the spec")
    check(len(table.rows) == 11, f"package table has {len(table.rows)} rows, expected 11")

    # Verticals keep pictures, never numerals.
    icons = [shp for shp in slides[2].shapes if shp.shape_type == 13]
    check(len(icons) == 4, f"{len(icons)} vertical icons, expected 4")

    # Architecture: named app, named engine product, one box per system.
    arch = texts(slides[7])
    check("Workforce optimization by SoftServe" in arch,
          "the architecture app box does not name the pack")
    check("NVIDIA cuOpt" in arch, "the engine box does not name the catalog product")
    check("BI" in arch, "the architecture has no destination box for the BI output")
    arrows = sum(1 for shp in slides[7].shapes if shp.shape_type == 9)
    check(arrows >= 4, f"{arrows} connectors on the architecture slide, expected 4+")


def run_variability(tmp: Path) -> None:
    out = tmp / "variability"
    rc, log = build(VARIABILITY, out)
    check(rc == 0, f"variability build exited {rc}, not 0:\n{log}")
    check("no overflow" in log, f"variability fit report is not clean:\n{log}")
    deck = out / "contract-intelligence-sales-deck.pptx"
    check(deck.is_file(), "variability deck was not written")
    prs = Presentation(str(deck))
    slides = list(prs.slides)
    check(len(slides) == 10, f"{len(slides)} slides, expected 10")

    # 3 verticals -> one greyed empty card, never a gap.
    icons = [shp for shp in slides[2].shapes if shp.shape_type == 13]
    check(len(icons) == 3, f"{len(icons)} vertical icons for 3 verticals")
    cards = [shp for shp in slides[2].shapes
             if shp.shape_type == 1 and round(shp.width / 914400, 2) == 2.25]
    check(len(cards) == 4, f"{len(cards)} vertical panels, expected 4 (the 4th greyed)")

    # 7 capability rows in both tables, cloned from the exemplar's own row.
    t9 = next(shp.table for shp in slides[8].shapes
              if getattr(shp, "has_table", False) and shp.has_table)
    t10 = next(shp.table for shp in slides[9].shapes
               if getattr(shp, "has_table", False) and shp.has_table)
    check(len(t9.rows) == 12, f"packages table has {len(t9.rows)} rows, expected 12")
    check(len(t10.rows) == 9, f"detailed table has {len(t10.rows)} rows, expected 9")

    # A 3-layer stack shortens the ladder; a 3-product engine names all three.
    arch = texts(slides[7])
    for product in ("NVIDIA NeMo Agent Toolkit", "NVIDIA NIM", "NVIDIA NeMo"):
        check(product in arch, f"the engine box does not name {product}")
    check("Obligation dashboard" in arch, "the destination system is missing")
    layers = texts(slides[6])
    check("Extraction engine" in layers and "Infrastructure" in layers,
          "the ladder was not filled from architecture.stack[]")
    # A resized ladder resizes what it holds: nothing pokes out of its row.
    rows = [sp for sp in slides[6].shapes if sp.shape_type == 1
            and sp.width / 914400 >= 9.0 and sp.height / 914400 >= 0.5]
    poking = []
    for sp in slides[6].shapes:
        if sp in rows or sp.height <= 0:
            continue
        cy, cx = sp.top + sp.height / 2, sp.left + sp.width / 2
        row = next((r for r in rows if r.top <= cy <= r.top + r.height
                    and r.left <= cx <= r.left + r.width), None)
        if row is None:
            continue
        over = max(row.top - sp.top, (sp.top + sp.height) - (row.top + row.height)) / 914400
        if over > 0.02:
            poking.append((sp.name, round(over, 2)))
    check(not poking, f"shapes stick out of their ladder row: {poking[:3]}")


def run_figureless(tmp: Path) -> None:
    """No cleared figure, `-` consumption, {label, text} problem points, a PoV cell that
    carries its own glyph, a destination named with a qualifier: the deck says so
    instead of printing dicts, dashes, figure caveats or the internal paragraph."""
    from deckkit import packspec_module
    packspec = packspec_module()
    spec = packspec.load(VARIABILITY)[0]
    for k in spec["kpis"]:
        k["figure"] = "-"
    spec["packages"]["target_oci_consumption"] = "-"
    spec["problem_solution"]["problem_points"] = [
        {"label": "Slow answers", "text": "a single obligation question takes days of reading"},
        {"label": "Missed renewals", "text": "auto-renewals pass unnoticed and lock in old terms"}]
    row0 = spec["packages"]["capability_handling"][0]
    row0["pov"] = "● " + row0["pov"]
    spec["architecture"]["outputs"][1]["system"] = "Any BI tool — Obligation dashboard included"
    spec["meta"]["source_engagement"]["divergence_from_pack"] = "INTERNAL PARAGRAPH never printed"
    spec["meta"]["source_engagement"]["divergence_line"] = "The pack generalizes the rules the proof of value hard-coded."
    path = tmp / "figureless-pack-spec.md"
    problems = packspec.roundtrip_problems(spec, str(path))
    check(not problems, f"the figure-less variant does not round-trip: {problems[:1]}")
    path.write_text(packspec.dump(spec), encoding="utf-8")
    out = tmp / "figureless"
    rc, log = build(path, out)
    check(rc == 0, f"figure-less build exited {rc}, not 0:\n{log}")
    deck = out / "contract-intelligence-sales-deck.pptx"
    check(deck.is_file(), "figure-less deck was not written")
    slides = list(Presentation(str(deck)).slides)
    t2 = texts(slides[1])
    check("{'label'" not in t2 and "Slow answers" in t2, "problem points print as raw dicts")
    t5 = texts(slides[4])
    check("First engagement" in t5 and "Figures from" not in t5,
          "the figure-less proof slide still attributes figures")
    check("illustrative" not in t5, "a figure caveat printed with no figures")
    check("INTERNAL PARAGRAPH" not in t5, "the proof slide prints the internal divergence paragraph")
    check("SOLUTION" in t5 and "VALUE FOR CLIENT" in t5,
          "the proof slide does not carry the exemplar's four blocks")
    check("measured in the proof of value" in t5,
          "a figure-less pack drops the stat strip instead of naming what is measured")
    t6 = texts(slides[5])
    check("OCI CONSUMPTION" in t6 and "To be defined" in t6,
          "a `-` consumption target does not keep the panel as an empty instance")
    check("indicative" not in t6, "the seller footnote mentions figures with none in the spec")
    t7 = texts(slides[6])
    check("TECHNOLOGY STACK" in t7 and "Tailored solution" in t7 and "Oracle AI accelerator" in t7,
          "the ladder slide does not carry the reference's title and cards")
    t8 = texts(slides[7])
    check("Any BI tool" in t8 and "Any BI tool — Obligation" not in t8,
          "a qualified destination is not split into name and second line")
    t9 = next(shp.table for shp in slides[8].shapes
              if getattr(shp, "has_table", False) and shp.has_table)
    first = len(t9.rows) - len(spec["packages"]["capability_handling"])
    check(t9.cell(first, 1).text.strip() == "●",
          f"PoV glyph carried in the prose is {t9.cell(first, 1).text.strip()!r}, not ●")
    t10 = next(shp.table for shp in slides[9].shapes
               if getattr(shp, "has_table", False) and shp.has_table)
    first10 = len(t10.rows) - len(spec["packages"]["capability_handling"])
    check("● ●" not in t10.cell(first10, 1).text and "●●" not in t10.cell(first10, 1).text,
          f"detailed PoV cell doubles the glyph: {t10.cell(first10, 1).text[:40]!r}")
    import re as _re
    ragged = [t10.cell(r, c).text[:24] for r in range(first10, len(t10.rows)) for c in (1, 2, 3)
              if t10.cell(r, c).text.strip() not in ("", "—")
              and not _re.match(r"^(●●|●|◐)  \S", t10.cell(r, c).text)]
    check(not ragged, f"detailed cells space their glyph unevenly: {ragged[:3]}")


def main() -> int:
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        run_fixture(tmp)
        run_variability(tmp)
        run_figureless(tmp)
    if failures:
        print(f"{len(failures)} of {checks} checks failed:")
        for f in failures:
            print("  - " + f)
        return 1
    print(f"{checks} checks passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
